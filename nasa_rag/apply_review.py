"""Apply verified RAG review edits to a COPY of the week3 workbook and log them in a review sheet.

usage: uv run --project ~/nasa_rag python apply_review.py <review_dir> <in.xlsx> <out.xlsx> [min_confidence]
"""
import glob, json, sys
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment

review_dir, src, dst = sys.argv[1:4]
min_conf = sys.argv[4] if len(sys.argv) > 4 else "high"
allowed = {"high"} if min_conf == "high" else {"high", "medium"}
skip = set(json.load(open(f"{review_dir}/skip.json"))) if glob.glob(f"{review_dir}/skip.json") else set()

wb = load_workbook(src)
edits, unver, citeprob, confirmed = [], [], [], []
for f in sorted(glob.glob(f"{review_dir}/out_*.json")):
    d = json.load(open(f))
    edits += [e for e in d.get("edits", [])]
    unver += d.get("unverifiable", [])
    citeprob += d.get("citation_problems", [])
    confirmed += d.get("confirmed_correct", [])

fill = PatternFill("solid", fgColor="FFF2CC")
applied, skipped = [], []
for e in edits:
    key = f"{e['sheet']}!{e['cell']}"
    if e.get("confidence", "medium") not in allowed or key in skip or e["sheet"] not in wb.sheetnames:
        skipped.append(e); continue
    ws = wb[e["sheet"]]
    cell = ws[e["cell"]]
    live = "" if cell.value is None else str(cell.value)
    if live.strip() != str(e["current"]).strip():
        e["note"] = "cell text differs from what the reviewer saw; not applied"
        skipped.append(e); continue
    cell.value = e["proposed"].replace("—", ",")
    cell.fill = fill
    cell.comment = Comment(f"RAG review 2026-08-28: {e['reason']} Source: {e['citation']}", "NB Tech RAG")
    applied.append(e)

# global: no em dashes anywhere
dash = 0
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and "—" in c.value:
                c.value = c.value.replace(" — ", ", ").replace("—", ","); dash += 1

if "RAG Review" in wb.sheetnames:
    del wb["RAG Review"]
ws = wb.create_sheet("RAG Review")
ws.append(["RAG review against the References library only, 2026-08-28. Edited cells are highlighted yellow and carry a comment with the source."])
ws.append([f"Applied {len(applied)} edits, {len(skipped)} proposed edits held back, {len(confirmed)} claims confirmed, {len(unver)} claims not coverable by the library, {dash} em dashes removed."])
ws.append([])
ws.append(["Status", "Sheet", "Cell", "Was", "Now / claim", "Reason", "Source", "Quote", "Confidence"])
for e in applied:
    ws.append(["EDITED", e["sheet"], e["cell"], e["current"], e["proposed"], e["reason"], e["citation"], e.get("quote", ""), e.get("confidence", "")])
for e in skipped:
    ws.append(["HELD (review manually)", e["sheet"], e["cell"], e["current"], e["proposed"], e.get("note", e["reason"]), e["citation"], e.get("quote", ""), e.get("confidence", "")])
for e in citeprob:
    ws.append(["CITATION PROBLEM", e["sheet"], e["cell"], "", e["problem"], "", e.get("citation", ""), "", ""])
for e in confirmed:
    ws.append(["CONFIRMED", e["sheet"], e["cell"], "", e["claim"], "", e.get("citation", ""), "", ""])
for e in unver:
    ws.append(["NOT IN LIBRARY", e["sheet"], e["cell"], "", e["claim"], e.get("why", ""), "", "", ""])
for c in ws[4]:
    c.font = Font(bold=True)
ws.freeze_panes = "A5"
for col, w in zip("ABCDEFGHI", (22, 20, 7, 40, 40, 40, 30, 40, 10)):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=5):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
wb.save(dst)
print(f"applied={len(applied)} held={len(skipped)} confirmed={len(confirmed)} unverifiable={len(unver)} citation_problems={len(citeprob)} em_dashes_fixed={dash}")
for e in applied:
    print(f"  {e['sheet']}!{e['cell']}: {e['current'][:60]!r} -> {e['proposed'][:60]!r}  {e['citation']}")
