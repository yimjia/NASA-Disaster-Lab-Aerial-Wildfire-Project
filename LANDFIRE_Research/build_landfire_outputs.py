"""LANDFIRE evaluation document (Word) and the Perimeters + LANDFIRE deck (PowerPoint), from the same numbers.

Inputs: landfire_summary.json and landfire_<fire>_<table>.csv (landfire_extract.py), figures/ (maps.py),
../Fireline_Operations/line_snapshots_daily.csv, ../Fuels_and_Topography/terrain_stats.json.
Run: uv run --with python-docx --with python-pptx python build_landfire_outputs.py
"""
import csv, json
from collections import defaultdict
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as PI, Pt as PP
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR

HERE = Path(__file__).parent
FIG = HERE / "figures"
S = json.load(open(HERE / "landfire_summary.json"))
TERR = json.load(open(HERE.parent / "Fuels_and_Topography" / "terrain_stats.json"))
SNAP = list(csv.DictReader(open(HERE.parent / "Fireline_Operations" / "line_snapshots_daily.csv")))
DATE = "22 October 2026"
KEYS = {"Pioneer 2016": "Pioneer", "Carr 2018": "Carr", "Camp 2018": "Camp", "LNU 2020": "LNU"}

# Scott and Burgan (2005) descriptions, condensed from the key and model descriptions (lookup/Scott_Burgan_2005_RMRS-GTR-153.pdf)
FM = {"TU5": "heavy conifer litter with shrub or small-tree understory", "TU1": "low load grass or shrub with litter",
      "TL1": "compact forest litter, light load", "TL3": "moderate load conifer litter", "TL4": "fine litter with small downed logs",
      "TL6": "moderate load broadleaf litter", "TL8": "moderate load long-needle pine litter", "SH2": "moderate load woody shrubs, about 1 ft",
      "SH5": "heavy shrub load, 4 to 6 ft", "SH7": "very heavy shrub load, 4 to 6 ft", "GS1": "shrubs about 1 ft, low grass load",
      "GS2": "shrubs 1 to 3 ft, moderate grass load", "GR2": "grass, low load"}
SB = {r["FBFM40"]: r for r in csv.DictReader(open(HERE / "lookup" / "scott_burgan_2005_table7.csv"))}


def load(code):
    r = SB[code]
    return sum(float(r[k]) for k in ("load_1h_tac", "load_10h_tac", "load_100h_tac", "load_live_herb_tac", "load_live_woody_tac"))


def snap(fire, day):
    r = [a for a in SNAP if a["fire"] == fire and a["snapshot_utc"].startswith(day)]
    return r[0] if r else None


def fccs_by_name(fire):
    agg = defaultdict(float)
    for r in csv.DictReader(open(HERE / f"landfire_{fire}_FCCS.csv")):
        agg[r["name"]] += float(r["share"])
    return sorted(agg.items(), key=lambda x: -x[1])


def pct(x):
    return f"{round(100 * x)}%"


OPS = {"Carr": ("2018-08-20", "20 Aug 2018"), "Camp": ("2018-11-23", "23 Nov 2018"), "LNU": ("2020-09-10", "10 Sep 2020")}
TER = {"Pioneer": "Pioneer 2016", "Carr": "Carr 2018", "Camp": "Camp 2018", "LNU": "LNU 2020 Hennessey"}

# What each picture shows: data-grounded readings of the maps and tables, written once and used in both outputs.
READ = {
    "Pioneer": {
        "fuels": ["LANDFIRE 2016 classifies this area after the fire: the disturbance layer marks 99% of it as burned in the prior 2 to 5 years, and 62% of the vegetation as 'Recently Burned-Herb and Grass Cover'.",
                  "So the fuel models here (TL3 conifer litter 23%, TL6 broadleaf litter 20%, grass and grass-shrub 26%) describe what was left, not what burned.",
                  "The dated perimeters fill the southern lobe first (light) and the northern lobe late (dark), consistent with the late-August northeastward run in the case-study record."],
        "second": ["The disturbance map is the evidence: almost every pixel inside the final perimeter carries a fire code, with high, moderate and low severity interleaved.",
                   "A pre-fire picture needs LANDFIRE 2014 (archived; the LANDFIRE viewer, not the image service).",
                   "No public fireline archive exists for 2016; incident GIS has to come through the Pioneer FOIA."]},
    "Carr": {
        "fuels": ["Half the area (49%) is TU5, heavy conifer litter under shrubs and small trees, the heaviest common model here at 14 t/ac assigned surface load.",
                  "Shrub (14%) and grass-shrub (13%) fill the lower slopes; Whiskeytown Lake is the non-burnable hole in the middle.",
                  "The 63 dated perimeters thicken in the south early and expand the north late; LANDFIRE marks 97% of the area as undisturbed in the 10 years before 2016, so this is a pre-fire picture."],
        "second": ["Canopy covers 81% of the area (median cover 55%, canopy base height median 1.6 m): low, continuous canopy that lets surface fire climb.",
                   "By 20 Aug the incident map held about 306 mi of completed dozer line and 30 mi of hand line, much of it cut through the interior as well as along the edge.",
                   "The incident map also carried 94 safety zones, 8 helispots and 6 unimproved landing areas (points on the map): ready-made checkpoints for an overhead drone."]},
    "Camp": {
        "fuels": ["Two landscapes: grass and grass-shrub (29% together) on the lower western slopes toward the valley, and timber with shrub understory (TU5, 36%) on the higher eastern ridges.",
                  "A straight north-south boundary in the middle, where timber litter switches to timber-understory, looks like a LANDFIRE map-zone seam, not a real vegetation edge.",
                  "LANDFIRE marks 98% of the area as undisturbed before 2016, so this is a pre-fire picture; the town of Paradise shows as timber and non-burnable (LANDFIRE codes developed land as non-burnable), so structure-to-structure spread is invisible in these fuels."],
        "second": ["Canopy covers 71% of the area (median cover 45%) and is densest in the east; the west is open grass.",
                   "By 23 Nov the incident map held about 145 mi of completed dozer line and 16 mi of hand line; dozer line is concentrated on the eastern and southern edges, where the fire met forest and canyons.",
                   "Few operations points are in the archive, and no line features at all from 13 to 18 Nov."]},
    "LNU": {
        "fuels": ["The complex is a mix: shrub (29%), timber-understory (22%), grass (21%) and grass-shrub (14%); Hennessey in the east is grass and shrub (chaparral), Walbridge in the west is timber.",
                  "Assigned surface load is median 8.3 t/ac, with the heavy shrub SH7 (17%) and TU5 (17%) driving the high end.",
                  "The 180 FEDS satellite perimeters (12 h) outline both fires; LANDFIRE marks 88% of the area undisturbed and about 9% as burned in the 2 to 5 years before 2016 (earlier fires inside the footprint)."],
        "second": ["Canopy covers 53% of the area: dense forest at Walbridge, open woodland and chaparral at Hennessey.",
                   "By 10 Sep the incident map held about 298 mi of completed dozer line, 25 mi of hand line and 239 mi of road used as line; uncontrolled edge had fallen from 486 mi (20 Aug) to zero.",
                   "Walbridge is ringed by dozer line; across the complex 239 mi of road served as containment line."]},
}

# ======================================================================== Word document
d = Document()
st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12); st.font.color.rgb = RGBColor(0, 0, 0)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman"); st.paragraph_format.space_after = Pt(4)
for s in d.sections:
    s.left_margin = s.right_margin = Inches(1); s.top_margin = s.bottom_margin = Inches(0.9)


def run(p, t, b=False, i=False, size=None):
    r = p.add_run(t); r.bold = b; r.italic = i; r.font.name = "Times New Roman"; r.font.color.rgb = RGBColor(0, 0, 0)
    if size: r.font.size = Pt(size)


def H(t, size=13):
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(8); p.paragraph_format.keep_with_next = True; run(p, t, b=True, size=size)


def P(t, i=False):
    p = d.add_paragraph(); run(p, t, i=i)


def B(t):
    p = d.add_paragraph(style="List Bullet"); run(p, t)


def T(hdr, rows, size=10):
    t = d.add_table(rows=1, cols=len(hdr)); t.style = "Table Grid"
    for j, h in enumerate(hdr):
        t.rows[0].cells[j].text = ""; run(t.rows[0].cells[j].paragraphs[0], h, b=True, size=size)
    for r in rows:
        c = t.add_row().cells
        for j, v in enumerate(r):
            c[j].text = ""; run(c[j].paragraphs[0], str(v), size=size)
    d.add_paragraph()


p = d.add_paragraph(); run(p, "LANDFIRE Evaluation of the Benchmark Fires", b=True, size=16)
P(f"NASA Aerial Wildfire Project, ACERO Disaster Lab working group. {DATE}.", i=True)
H("Purpose")
P("Document what LANDFIRE says about the fuels the four benchmark fires burned through (Pioneer 2016, Carr 2018, Camp 2018, LNU Lightning Complex 2020), "
  "how the data were evaluated, what they can and cannot support, and how they line up with the fire perimeters and the fireline built on each fire.")
H("What was evaluated")
B("LANDFIRE product versions (landfire.gov comparison table): LF 2014, LF 2016 Remap, LF 2019 Limited, LF 2020. LF 2016 Remap is the newest version that predates Carr, Camp and LNU, and is the only pre-2018 version on the public LANDFIRE image services.")
B("Layers used, all 30 m from LF 2016 Remap: FBFM40 (Scott and Burgan 40 fuel models), FCCS fuelbeds, EVT (existing vegetation type), canopy cover (CC, %), canopy bulk density (CBD, kg m-3 x 100), canopy base height (CBH, m x 10), canopy height (CH, m x 10), fuel disturbance (FDist). Units from the LANDFIRE Data Dictionary.")
B("Code tables from landfire.gov (LF16_F40_200, LF16_FCCS_200, LF16_EVT_200, LF16_FDst_200). Surface fuel load per pixel from Scott and Burgan (2005), Table 7: the model's assigned load (1-h, 10-h, 100-h, live herb, live woody), not a measurement.")
B("Accuracy evidence: LANDFIRE-type fuel classifications explained little of plot-measured surface loading (R2 below 0.30 for most components; Keane et al., 2013); canopy fuel layers are rated low accuracy at national scale (LANDFIRE canopy fuels review). These limits apply to every number below.")
B("Considered and not used: WFIGS Daily Perimeters (NIFC item 7fa2437e625d49f7af1017c8617b68c1). It keeps every geometry change of WFIGS fire perimeters, ordered by burn period, but its earliest record is 4 May 2021, so it holds none of the four benchmark fires (2016 to 2020). It is the right source for fires from mid-2021 on, such as Eaton 2025. For our fires the dated perimeters come from NIFC GeoMAC history and NASA FEDS, and fireline from the NIFC Operational Data Archive 2018 and 2020.")
H("Method")
P("For each fire the layers were requested from the LANDFIRE LF2016 image services (lfps.usgs.gov) over the final perimeter's bounding box at 30 m in EPSG:5070, "
  "in tiles where the area was large, and only pixels inside the final perimeter were kept. Shares are of the burned area. Scripts: landfire_extract.py (numbers), maps.py (figures). "
  "Perimeters: project copies of the official final perimeters; dated perimeters from NIFC GeoMAC history (Pioneer, Carr, Camp) and NASA FEDS 12 h satellite perimeters (LNU). Fireline: NIFC Operational Data Archive 2018 and 2020 (Research/Fireline_Operations).")
H("Results across the four fires")
rows = []
for k, v in S.items():
    g = v["fuel_group_share"]
    rows.append([k, f"{v['acres']:,}", pct(g.get("Grass", 0) + g.get("Grass-shrub", 0)), pct(g.get("Shrub", 0)),
                 pct(g.get("Timber-understory", 0) + g.get("Timber litter", 0)), f"{v['surface_load_tac_burnable']['median']} ({v['surface_load_tac_burnable']['p90']})",
                 pct(v["canopy_share_of_area"]), f"{v['canopy_cover_pct_where_canopy']['median']:.0f}", f"{v['cbh_m']['median']}", f"{v['cbd_kg_m3']['median']}"])
T(["Fire", "Acres (pixels)", "Grass + grass-shrub", "Shrub", "Timber", "Surface load t/ac, median (p90)", "Area with canopy", "Canopy cover % (median)", "CBH m (median)", "CBD kg/m3 (median)"], rows, size=9)
P("Timber dominates Carr and Camp, shrub and grass dominate LNU, and Pioneer cannot be read as pre-fire (below). The highest assigned surface loads (Carr median 14 t/ac) come from the TU5 fuel model, heavy conifer litter with a shrub understory.")
for k, v in S.items():
    key = KEYS[k]
    H(k, size=12)
    fm = "; ".join(f"{c} {pct(s_)} ({FM.get(c, 'see Scott and Burgan 2005')}, {load(c):.1f} t/ac)" for c, s_ in v["top_fuel_models"][:5])
    P(f"Top fuel models: {fm}.")
    P("Top FCCS fuelbeds: " + "; ".join(f"{n} {pct(s_)}" for n, s_ in fccs_by_name(key)[:4]) + ".")
    P("Top vegetation types (EVT): " + "; ".join(f"{n} {pct(s_)}" for n, s_ in v["top_evt"][:4]) + ".")
    P(f"Canopy: {pct(v['canopy_share_of_area'])} of the area has canopy; cover median {v['canopy_cover_pct_where_canopy']['median']:.0f}% "
      f"(p10 {v['canopy_cover_pct_where_canopy']['p10']:.0f}, p90 {v['canopy_cover_pct_where_canopy']['p90']:.0f}); canopy height median {v['ch_m']['median']} m; "
      f"base height median {v['cbh_m']['median']} m; bulk density median {v['cbd_kg_m3']['median']} kg/m3.")
    P("Disturbance before LF 2016: " + "; ".join(f"{n} {pct(s_)}" for n, s_ in v["fdst_top"][:3]) + ".")
    t = TERR.get(TER[key])
    if t:
        P(f"Terrain (USGS 3DEP, Research/Fuels_and_Topography): median elevation {t['elev_m']['median']} m, median slope {t['slope_deg']['median']}°, {pct(t['slope_share']['30deg_plus'])} of the area at 30° or more.")
    for line in READ[key]["fuels"] + READ[key]["second"]:
        B(line)
H("Data quality findings")
B("Pioneer: LF 2016 Remap already contains the Pioneer Fire. FDist marks 99% of the area as fire-disturbed in the prior 2 to 5 years and EVT classes 62% as recently burned. Its fuel models are post-fire. Use LF 2014 for a pre-fire picture.")
B("FCCS value 0 is labelled 'Bare Ground' in LANDFIRE's table, yet 56% (Camp) and 62% (LNU) of those pixels carry grass fuel models. Do not read FCCS 0 as bare ground in these landscapes.")
B("Camp shows a straight north-south boundary where timber litter switches to timber-understory, consistent with a LANDFIRE map-zone seam (a documented artifact class), not a vegetation edge.")
B("Surface fuel loads here are the loads Scott and Burgan assign to each fuel model, not measured loads; plot-level loading is poorly predicted by fuel classifications (Keane et al., 2013).")
B("Canopy base height p90 sits at 10 m on every fire, which suggests a cap in the product; treat the upper tail as censored.")
B("LF 2016 Remap predates Carr and Camp by about two years and LNU by about four; recent disturbance after 2016 (for example the 2017 and 2018 fires in the LNU footprint) is not reflected.")
B("In the fireline archive, the generic 'Completed Line' category is inflated by re-posts (918 mi at Carr) and is not used; dozer and hand line are reported from daily snapshots.")
H("Next steps")
B("Download LF 2014 FBFM40, canopy and EVT for Pioneer through the LANDFIRE viewer and rerun landfire_extract.py on it.")
B("Compare LF 2016 Remap with LF 2019 Limited for LNU to see what changed after the 2017 and 2018 fires inside its footprint.")
B("Overlay daily fireline snapshots on fuel groups to measure how much line was built in each fuel type (where crews and dozers worked).")
B("Pair FBFM40 with the RAWS and gridMET fuel moisture already in the project to run fire behavior per fuel model on each peak-run day.")
H("References")
for r in ["Keane, R. E., Herynk, J. M., Toney, C., Urbanski, S. P., Lutes, D. C., & Ottmar, R. D. (2013). Evaluating the performance and mapping of three fuel classification systems using FIA surface fuel measurements. Forest Ecology and Management, 305, 248-263. https://doi.org/10.1016/j.foreco.2013.06.001",
          "LANDFIRE. (n.d.). LANDFIRE product versions comparison table; LF 2016 Remap; LANDFIRE Data Dictionary; LF2016 attribute tables (LF16_F40_200, LF16_FCCS_200, LF16_EVT_200, LF16_FDst_200). U.S. Department of the Interior, Geological Survey, and U.S. Department of Agriculture. https://www.landfire.gov",
          "Scott, J. H., & Burgan, R. E. (2005). Standard fire behavior fuel models: A comprehensive set for use with Rothermel's surface fire spread model (Gen. Tech. Rep. RMRS-GTR-153). USDA Forest Service, Rocky Mountain Research Station.",
          "National Interagency Fire Center. (2018, 2020). Operational Data Archive (National Incident Feature Service). https://data-nifc.opendata.arcgis.com"]:
    p = d.add_paragraph(); p.paragraph_format.left_indent = Inches(0.5); p.paragraph_format.first_line_indent = Inches(-0.5); run(p, r)
for p in d.paragraphs:
    for r in p.runs:
        r.font.name = "Times New Roman"; r.font.color.rgb = RGBColor(0, 0, 0)
d.save(HERE / "LANDFIRE_Evaluation_Benchmark_Fires_2026-10-22.docx")
print("wrote docx")

# ======================================================================== PowerPoint (v0.5 deck look)
C = {k: PRGB.from_string(v) for k, v in dict(ink="1D2125", ink2="3A424B", bg="FFFFFF", bg2="F1F2F4", ember="E8642F", slate="6B7480",
                                             amber="E9A821", green="18A878", grid="E1E4E8").items()}
W, M = 13.333, 0.55
prs = Presentation(); prs.slide_width = PI(W); prs.slide_height = PI(7.5); BL = prs.slide_layouts[6]
FOOT = f"NASA Aerial Wildfire Project  |  ACERO Disaster Lab  |  Perimeters and LANDFIRE  |  {DATE}"


def tx(s, x, y, w, h, paras, size=14, color=None, bold=False, font="Calibri", anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(PI(x), PI(y), PI(w), PI(h)); tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for a in ("margin_left", "margin_right", "margin_top", "margin_bottom"): setattr(tf, a, 0)
    for i, para in enumerate(paras if isinstance(paras, list) else [paras]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.space_after = PP(6)
        for seg in (para if isinstance(para, list) else [para]):
            t, o = (seg, {}) if isinstance(seg, str) else seg
            r = p.add_run(); r.text = t; r.font.size = PP(o.get("size", size)); r.font.bold = o.get("bold", bold)
            r.font.name = o.get("font", font); r.font.color.rgb = o.get("color", color or C["ink"])


def box(s, x, y, w, h, fill):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PI(x), PI(y), PI(w), PI(h)); sh.adjustments[0] = 0.05
    sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.fill.background(); sh.shadow.inherit = False


n = [0]


def slide(kicker, title):
    n[0] += 1
    s = prs.slides.add_slide(BL)
    tx(s, M, 0.32, 10, 0.3, kicker.upper(), size=12, bold=True, color=C["ember"], font="Arial")
    tx(s, M, 0.6, 12.2, 0.7, title, size=26, bold=True, font="Arial")
    tx(s, M, 7.08, 10, 0.25, FOOT, size=9, color=C["slate"]); tx(s, 12.3, 7.08, 0.5, 0.25, str(n[0] + 1), size=9, color=C["slate"])
    return s


def picture_slide(kicker, title, img, heading, bullets, note=None):
    s = slide(kicker, title)
    from PIL import Image
    iw, ih = Image.open(img).size
    maxw, maxh = 8.1, 5.6
    sc = min(maxw / iw, maxh / ih); w, h = iw * sc, ih * sc
    s.shapes.add_picture(str(img), PI(M), PI(1.4 + (maxh - h) / 2), PI(w), PI(h))
    x = M + maxw + 0.25; bw = W - M - x
    box(s, x, 1.4, bw, 5.55, C["bg2"])
    tx(s, x + 0.2, 1.52, bw - 0.4, 0.35, heading, size=15, bold=True, font="Arial")
    tx(s, x + 0.2, 1.95, bw - 0.4, 4.9, [[("•  ", {"bold": True, "color": C["ember"]}), b] for b in bullets] +
       ([[(note, {"size": 10, "color": C["slate"]})]] if note else []), size=12)
    return s


# title
s = prs.slides.add_slide(BL)
box(s, -0.1, -0.1, W + 0.2, 7.7, C["ink"])
tx(s, M, 2.0, 12, 1.4, "What the Fuels and Perimeters Show", size=40, bold=True, color=C["bg"], font="Arial", anchor=MSO_ANCHOR.BOTTOM)
tx(s, M, 3.6, 11.5, 1.2, "LANDFIRE fuels, canopy and disturbance under the dated perimeters and the fireline built on Pioneer 2016, Carr 2018, Camp 2018 and the LNU Lightning Complex 2020",
   size=19, color=C["bg2"])
tx(s, M, 6.6, 12, 0.4, f"NASA Aerial Wildfire Project  |  ACERO Disaster Lab working group  |  {DATE}  |  For evaluation: what each dataset displays and what it can support", size=12, color=C["slate"])

# how to read
s = slide("How to read these slides", "Three datasets, layered on the same map")
cards = [("LANDFIRE 2016 Remap", "30 m maps of fuel model (Scott and Burgan 40), canopy cover, height, base height, bulk density, vegetation type and recent disturbance. "
          "Circa-2016 conditions: pre-fire for Carr, Camp and LNU; post-fire for Pioneer."),
         ("Dated perimeters", "NIFC GeoMAC perimeter history for Pioneer (89), Carr (63) and Camp (28); NASA FEDS 12 h satellite perimeters for LNU (180). "
          "Light lines are early, dark lines are late. No GOFER."),
         ("Fireline and operations", "NIFC Operational Data Archive: completed dozer and hand line, road used as line, safety zones, helispots and landing areas, "
          "as they stood on the incident map on one day near the end of each fire.")]
for i, (h, b) in enumerate(cards):
    x = M + i * 4.12
    box(s, x, 1.5, 3.95, 3.3, C["bg2"]); tx(s, x + 0.2, 1.65, 3.55, 0.4, h, size=16, bold=True, font="Arial"); tx(s, x + 0.2, 2.15, 3.55, 2.6, b, size=13, color=C["ink2"])
box(s, M, 5.0, 12.2, 1.85, C["ink"])
tx(s, M + 0.25, 5.15, 11.7, 1.6, [[("Read with care: ", {"bold": True, "color": C["amber"]}), "fuel loads are what each fuel model assigns, not measurements, and fuel classifications explain little of measured loading (R2 below 0.30; Keane et al., 2013). "
      "Pioneer's LANDFIRE map is after the fire. Fireline dates are when lines were mapped, not when they were cut."]], size=14, color=C["bg"])

# comparison
s = slide("Across the four fires", "Timber at Carr and Camp, shrub and grass at LNU")
s.shapes.add_picture(str(FIG / "compare_fuel_groups.png"), PI(M), PI(1.35), PI(8.6))
rows = [("Fire", "Canopy area", "Cover", "CBH", "Load")] + [(k.split()[0], pct(v["canopy_share_of_area"]), f"{v['canopy_cover_pct_where_canopy']['median']:.0f}%",
                                                              f"{v['cbh_m']['median']} m", f"{v['surface_load_tac_burnable']['median']} t/ac") for k, v in S.items()]
tbl = s.shapes.add_table(len(rows), 5, PI(9.4), PI(1.5), PI(3.4), PI(2.6)).table
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        c = tbl.cell(i, j); c.text = v; p_ = c.text_frame.paragraphs[0]; p_.runs[0].font.size = PP(11); p_.runs[0].font.bold = i == 0
        c.fill.solid(); c.fill.fore_color.rgb = C["ink"] if i == 0 else (C["bg2"] if i % 2 else C["bg"])
        p_.runs[0].font.color.rgb = C["bg"] if i == 0 else C["ink"]
tx(s, 9.4, 4.3, 3.4, 2.6, ["Carr carries the heaviest assigned surface load (median 14 t/ac) because half its area is TU5, conifer litter under shrubs.",
                           "Canopy base heights of 1 to 2 m on the California fires mean low canopy that surface fire can climb.",
                           [("Pioneer values are post-fire.", {"bold": True})]], size=12)

for k, v in S.items():
    key = KEYS[k]
    picture_slide(f"{k}  |  fuels and perimeters", {"Pioneer": "Pioneer: LANDFIRE shows what the fire left, not what it burned",
                                                    "Carr": "Carr: heavy timber litter under shrubs, nearly all pre-fire",
                                                    "Camp": "Camp: grass to the west, timber to the east",
                                                    "LNU": "LNU: chaparral and grass at Hennessey, timber at Walbridge"}[key],
                  FIG / f"{key}_fuels_progression.png", "What the picture shows", READ[key]["fuels"],
                  note="Fuel groups from LANDFIRE LF2016 FBFM40. Perimeter colors run light (early) to dark (late).")
    if key == "Pioneer":
        picture_slide(f"{k}  |  disturbance", "Pioneer: the 2016 map already records the burn", FIG / f"{key}_canopy_operations.png",
                      "What the picture shows", READ[key]["second"], note="LANDFIRE LF2016 fuel disturbance (FDist), fire severity classes.")
    else:
        day = OPS[key][1]
        picture_slide(f"{k}  |  canopy and fireline", {"Carr": "Carr: dense low canopy and about 306 miles of dozer line",
                                                       "Camp": "Camp: line built where the fire met forest and canyons",
                                                       "LNU": "LNU: dozer line rings Walbridge; roads served as line"}[key],
                      FIG / f"{key}_canopy_operations.png", "What the picture shows", READ[key]["second"],
                      note=f"Canopy cover from LANDFIRE LF2016. Lines and points from the NIFC Operational Data Archive, as on the incident map on {day}.")

s = slide("What the picture paints", "Fuels set where fire could go; the record shows where people worked")
pts = [("Two kinds of landscape", "Timber with a shrub understory (Carr, Camp east, Walbridge) and grass, grass-shrub and chaparral (Camp west, Hennessey). The California fires sit on low canopy over heavy litter."),
       ("Containment followed access", "Completed line hugs roads, ridges and the forest edge: about 306 mi of dozer line at Carr, 145 mi at Camp, 298 mi plus 239 mi of road at LNU."),
       ("The data have different clocks", "LANDFIRE is circa 2016, perimeters are hours to a day apart, fireline is when it was mapped. Every comparison has to name its time."),
       ("Where drones fit", "Fuel maps say which flanks run fast (grass, shrub); canopy and slope say where ground crews and line are slow; safety zones and landing areas are already mapped and can be checked from the air.")]
for i, (h, b) in enumerate(pts):
    x = M + (i % 2) * 6.15; y = 1.45 + (i // 2) * 2.75
    box(s, x, y, 6.0, 2.55, C["bg2"]); tx(s, x + 0.2, y + 0.15, 5.6, 0.4, h, size=16, bold=True, font="Arial"); tx(s, x + 0.2, y + 0.65, 5.6, 1.8, b, size=13, color=C["ink2"])

s = slide("Data quality", "What these pictures cannot tell us yet")
qs = [("Pioneer is post-fire", "LF 2016 marks 99% of the area fire-disturbed. Need LF 2014."),
      ("Loads are assigned", "Scott and Burgan model loads, not field measurements; classification R2 below 0.30."),
      ("FCCS 'Bare Ground'", "Most of those pixels carry grass fuel models (56% Camp, 62% LNU)."),
      ("Map-zone seam", "Straight north-south fuel boundary across Camp is an artifact."),
      ("Canopy base height cap", "p90 is 10 m on every fire: the upper tail is censored."),
      ("Fireline re-posts", "Daily re-posts inflate totals; we use one-day snapshots. No archive before 2018.")]
for i, (h, b) in enumerate(qs):
    x = M + (i % 3) * 4.1; y = 1.45 + (i // 3) * 2.75
    box(s, x, y, 3.95, 2.55, C["bg2"]); tx(s, x + 0.2, y + 0.15, 3.55, 0.4, h, size=15, bold=True, font="Arial", color=C["ember"]); tx(s, x + 0.2, y + 0.65, 3.55, 1.8, b, size=13, color=C["ink2"])

s = slide("Next steps", "Turning the pictures into measurements")
tx(s, M, 1.5, 12, 5, [[("1.  ", {"bold": True, "color": C["ember"]}), "Download LANDFIRE 2014 for Pioneer (LANDFIRE viewer) and redo its fuels as pre-fire."],
                      [("2.  ", {"bold": True, "color": C["ember"]}), "Measure miles of fireline built in each fuel group and slope class, per day."],
                      [("3.  ", {"bold": True, "color": C["ember"]}), "Compare LF 2016 with LF 2019 Limited at LNU to capture the 2017 and 2018 fires inside its footprint."],
                      [("4.  ", {"bold": True, "color": C["ember"]}), "Run fire behavior per fuel model with the peak-day fuel moisture already in the project (RAWS, gridMET)."],
                      [("5.  ", {"bold": True, "color": C["ember"]}), "Request Pioneer incident GIS (fireline, IAP maps) through the FOIA."]], size=18)
prs.save(HERE / "Perimeters_and_LANDFIRE_2026-10-22.pptx")
print("wrote pptx", len(prs.slides), "slides")
