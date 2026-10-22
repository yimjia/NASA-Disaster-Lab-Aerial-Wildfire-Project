"""Build the five mission uncertainty documents, version 1.1 (M1 to M5), plus the numbers the slide deck uses.

Changes from v1.0 (20 Oct 2026 review):
- Every number comes from data inside the project folder: Weather Data (weather_evaluation.py), the perimeter
  time-area records in perimeter_supplements/ (mission_timeline_evaluation.py, GOFER rows dropped), the Week 6
  workbook, and UQ proposal v0.5 (its Week 5 and library citations). Values we set ourselves are labelled "assumed".
- Night content restored from that data: night wind gate, share of aircraft perimeters flown at night, share of
  growth between aircraft perimeters that fell at night, and the night capability statements already in v0.5.
- Each document carries an hour-by-hour timeline from the perimeter records showing where drones fit: initial
  attack, deciding the perimeter, where resources go, escape routes, safety zones, helicopter landing zones.
- Meeting-note items that still need sources (crew GPS apps, fuel layers, satellite archives) are left out.

Run after weather_evaluation.py and mission_timeline_evaluation.py:
    uv run --with python-docx --with pandas python build_mission_uq_v11.py
"""
import csv, json, math, random, shutil
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn

HERE = Path(__file__).parent
OUT = HERE / "Mission UQ v1"; OUT.mkdir(exist_ok=True)
OLD = OUT / "superseded v1.0"; OLD.mkdir(exist_ok=True)
for f in OUT.glob("*_UQ_v1.0.docx"):
    shutil.move(str(f), OLD / f.name)
DATE = "20 October 2026"
VER = "1.1"
WX = json.load(open(HERE / "weather_eval" / "results.json"))
NIGHT = json.load(open(HERE / "timeline_eval" / "night_stats.json"))
BLACK = RGBColor(0, 0, 0)
REP = {"Camp": "JBGC1", "Carr": "MMOC1", "Pioneer": "BCFI1", "LNU": "ATLC1"}
AX, SB, T30, T40 = "Alta X (22 mph)", "SwitchBlade-Elite (25 mph)", "Type 1 FW if 30 mph", "Type 1 FW if 40 mph"
NUM = {}  # numbers for the deck


def tl(name):
    return list(csv.DictReader(open(HERE / "timeline_eval" / f"{name}_hourly.csv")))


def g(fire, drone, kind="crit", gate="p_fly_gust"):
    key = {"crit": "crit_gate", "perim": "perimeter_gate", "night": "perimeter_night_gate"}[kind]
    return WX[fire][REP[fire]][key][drone][gate]


def pct(x):
    return "under 1%" if 0 < x < 0.005 else f"{round(100 * x):d}%"


def age_h(R, p):
    """Expected information age (h): passes every R hours, each succeeds with p; E[A] = R (2 - p) / (2 p)."""
    return float("inf") if p == 0 else R * (2 - p) / (2 * p)


def max_age(rows, col):
    ages = [float(r[col].split(",")[-1].split("h")[0]) for r in rows if "h old" in r[col]]
    return max(ages) if ages else None


# ============================================================ document helpers
def new_doc():
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"; st.font.size = Pt(10); st.font.color.rgb = BLACK
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.space_after = Pt(3); st.paragraph_format.space_before = Pt(0)
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(0.7); s.top_margin = s.bottom_margin = Inches(0.5)
    return d


def run(p, text, bold=False, italic=False, size=None):
    r = p.add_run(text); r.bold = bold; r.italic = italic
    r.font.name = "Times New Roman"; r.font.color.rgb = BLACK
    if size: r.font.size = Pt(size)
    return r


def H(d, text, size=12):
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run(p, text, bold=True, size=size)


def P(d, *parts, after=3):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    for x in parts:
        if isinstance(x, tuple):
            run(p, x[0], bold="b" in x[1], italic="i" in x[1])
        else:
            run(p, x)
    return p


def B(d, lead, text):
    p = d.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(1)
    run(p, lead, bold=True); run(p, text)


def T(d, header, rows, widths, size=9):
    t = d.add_table(rows=1, cols=len(header)); t.style = "Table Grid"; t.autofit = False
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""
        run(c.paragraphs[0], h, bold=True, size=size)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            run(cells[i].paragraphs[0], str(v), size=size)
    from docx.oxml import OxmlElement
    mar = OxmlElement("w:tblCellMar")  # tight cell padding keeps the hourly tables on one page
    for side, v in (("top", 0), ("bottom", 0), ("left", 60), ("right", 60)):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
    t._tbl.tblPr.append(mar)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):  # Pages reads the grid
        gc.set(qn("w:w"), str(int(w * 1440)))
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
            for para in row.cells[i].paragraphs:
                para.paragraph_format.space_after = Pt(0)
    d.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def header_block(d, mission, title, note):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(1)
    run(p, f"{mission} {title}: Quantified Uncertainty, Version {VER}", bold=True, size=14)
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    run(p, f"NASA Aerial Wildfire Project, ACERO Disaster Lab working group. {DATE}. {note}", italic=True, size=11)


def timeline_table(d, rows):
    body, last_task, last_day = [], None, None
    short = lambda x: "none" if x == "none yet" else x.replace(" ac, ", " (").replace(" h old", " h)")
    for r in rows:
        day, hh = r["hour"][:6].lstrip("0"), r["hour"][-5:-3]
        label = (f"{day} {hh}" if day != last_day else hh) + (" N" if r["period"] == "Night" else "")
        last_day = day
        wind = "n/a" if not r["sust"] else f"{r['sust']}/{r['gust']}, {r['rh']}%"
        fly = "n/a" if not r["alta_x"] else f"{r['alta_x']}/{r['type1_30']}"
        task = r["task"] if r["task"] != last_task else ""
        last_task = r["task"]
        body.append([label, short(r["aircraft_perimeter"]), short(r["ics209"]), wind, fly, r["events"], task])
    T(d, ["Hour (N = night)", "Aircraft perimeter (age)", "ICS-209 (age)", "Wind mph sust/gust, RH", "Alta X / T1 30", "What the records show", "Drone task from this hour"],
      body, [0.62, 0.92, 0.85, 0.72, 0.52, 1.92, 1.65], size=8)


def night_line(fire):
    n = NIGHT[fire]
    if not n["aircraft_perimeters"]:
        return f"{fire} has no aircraft perimeter series in the record, so night growth is not computed without GOFER."
    return (f"On {fire}, {pct(n['night_stamped_share'])} of the {n['aircraft_perimeters']} aircraft perimeters were stamped between 20:00 and 06:00, "
            f"while about {pct(n['night_growth_share'])} of the growth between perimeters fell in those hours (growth spread evenly between perimeters).")


NIGHT_CAP = ("Night capability in project records: drones have mapped perimeters at night (ConOps 2026, p.14); the FVR-90 operator states day and night operation (W5); "
             "ACERO TCL 1 flights were daytime, within visual line of sight (Wu et al. 2025, p.9); night flight needs a DOI/FAA agreement (NWCG 2025, pp.6-7).")
SRC = ("Sources, all in the project folder: Weather Data (RAWS 10-min mean at 20 ft, gust = prior-hour peak); perimeter_supplements time-area records, GOFER rows dropped; "
       "Week 6 workbook; UQ proposal v0.5 (W5 = Week 5). Drone tasks and values marked assumed are our proposals, not records.")

# ============================================================ M1 Fireline mapping
d = new_doc()
header_block(d, "M1", "Fireline Mapping", "Version 1.1. Standard day for every mission and fire: 24 h, 06:00 to 06:00 local, on the peak-run day. GOFER not used.")
rows = tl("M1_Carr")
CAMP_G = 55000 / 12  # Green Sheet in the Camp record: more than 55,000 ac in the first 12 h
pA, p30 = g("Camp", AX), g("Camp", T30)
a_air, a_209 = max_age(rows, "aircraft_perimeter"), max_age(rows, "ics209")
NUM["M1"] = dict(a_air=a_air, a_209=a_209, pA=pA, p30=p30, ageS=age_h(1, g("Camp", AX, gate="p_fly_sustained")), ageA=age_h(1, pA), age30=age_h(1, p30), unseen30=CAMP_G * age_h(1, p30))
H(d, "Bottom line")
B(d, f"Up to {a_air:.0f} h old. ", f"During Carr's 26 July run the newest aircraft perimeter reached {a_air:.0f} h old and the newest ICS-209 {a_209:.0f} h, "
  f"yet station wind allowed every listed drone to fly all {len(rows)} hours. A 1 h drone revisit would have held map age near 0.5 h.")
B(d, "Camp is the hard case. ", f"Read against gusts, an Alta X could fly about {round(pA * 24)} hour in 24 on Camp's standard day ({pct(pA)}), so its map would be most of the day old. "
  f"A Type 1 fixed wing rated to 30 mph flies {pct(p30)} of hours: expected map age {age_h(1, p30):.1f} h, about {round(CAMP_G * age_h(1, p30), -2):,.0f} ac unmapped at any moment.")
B(d, "The largest unknowns are definitions, not weather. ", "Whether a drone's wind limit applies to sustained wind or gusts, and the unpublished Type 1 limit, move the answer more than any other input.")
H(d, "Who uses it")
P(d, ("Incident commander and planning (situation unit, fire behavior analyst): ", "b"), "map age decides whether resource orders and evacuation triggers rest on a current line. ",
  ("Division supervisors and lookouts: ", "b"), "head and flank position on the active edge, refreshed within the hour.")
H(d, "Drone types serving M1")
T(d, ["Drone", "Type", "Parameters", "Wind limit", "M1 role"], [
    ["L3Harris FVR-90", "NWCG Type 1 FW, VTOL", "12 to 18 h with 8 to 22 lb thermal video; Type 1 class 3,500 to 8,000 ft AGL, 50 mi", "Not found (30 and 40 mph run)", "Whole-fire perimeter, day and night"],
    ["SuperVolo (NASA Langley)", "VTOL FW research", "8 h; 45 to 75 mph; flown as relay in TCL 1", "Not found", "Second mapping tier, relay"],
    ["Alta X (Freefly)", "NWCG Type 3 rotor", "50 min empty, 22 min at 22 lb; FY22 average sortie 17 min", "22 mph (secondary)", "Head and spot detail close in"],
    ["SwitchBlade-Elite", "Type 3 band", "Up to 50 min, average sortie 18 min; 2 to 2.3 kg", "25 mph (W5)", "Close-in edge mapping"],
], [1.25, 1.1, 2.6, 1.0, 1.35], size=9)
H(d, "How the uncertainty is quantified")
P(d, ("Map error E(t) = G x A(t): ", "b"), "area grown since the newest map. With a pass every R hours that succeeds with probability p (the hourly wind gate), ",
  ("E[A] = R (2 - p) / (2p).", "b"), " Sustained wind gives the optimistic p, gusts the conservative p; both are 20 ft winds, so both overstate availability at flight altitude.")
res = []
for name, lab in [(AX, "Alta X, 22 mph"), (SB, "SwitchBlade-Elite, 25 mph"), (T30, "Type 1 if 30 mph"), (T40, "Type 1 if 40 mph")]:
    ps, pg = g("Camp", name, gate="p_fly_sustained"), g("Camp", name)
    res.append([lab, pct(ps), pct(pg), f"{age_h(1, ps):.1f} to {age_h(1, pg):.1f}", f"{round(CAMP_G * age_h(1, ps), -2):,.0f} to {round(CAMP_G * age_h(1, pg), -2):,.0f}",
                pct(g("Camp", name, "night"))])
T(d, ["Camp 8 to 9 Nov, R = 1 h", "p sustained", "p gust", "E[A] h", "Unseen ac", "p gust, Camp nights"], res, [1.6, 0.85, 0.7, 0.9, 1.4, 1.3])
P(d, ("Kinds of uncertainty: ", "b"), "p aleatory (weather); wind-limit definition and Type 1 limit epistemic; perimeter area has measurement error (the archive drops from 11,668 ac at 16:34 to 7,016 ac at 19:38 on 25 July, hence the noon rows). "
  "E[A] above 24 h or unseen area above 55,000 ac exceeds the standard day and the fire's own growth: read as no useful map. ",
  ("v1.2 needs: ", "b"), "Yiming's hourly perimeters as the reference for G; the FVR-90 wind limit; flight-level wind.", after=2)
H(d, "Night")
P(d, night_line("Carr"), " ", night_line("Camp"), f" Night wind at the stations allowed an Alta X in {pct(g('Carr', AX, 'night'))} of Carr night hours and {pct(g('Camp', AX, 'night'))} of Camp's (gusts). ", NIGHT_CAP)
H(d, "Hour by hour, standard day: Carr, 26 July 06:00 to 27 July 06:00 PDT (Mule Mountain RAWS)")
timeline_table(d, rows)
P(d, (SRC, "i"))
d.save(OUT / f"M1_Fireline_Mapping_UQ_v{VER}.docx")

# ============================================================ M2 Spot fire detection
d = new_doc()
header_block(d, "M2", "Spot Fire Detection", "Version 1.1. Standard day for every mission and fire: 24 h, 06:00 to 06:00 local, on the peak-run day. GOFER not used.")
rows = tl("M2_Camp")
R2 = 30


def tdet(R, p, Pd):
    q = p * Pd
    return float("inf") if q == 0 else R / 2 + R * (1 - q) / q


def gsd(h, hfov=40, n=640):
    return 2 * h * math.tan(math.radians(hfov / 2)) / n


NUM["M2"] = dict(best=tdet(R2, g("Camp", AX, gate="p_fly_sustained"), 0.9), worstA=tdet(R2, g("Camp", AX), 0.9), sb=tdet(R2, g("Camp", SB), 0.9), t30=tdet(R2, g("Camp", T30), 0.9))
H(d, "Bottom line")
B(d, f"{NUM['M2']['best']:.0f} to {NUM['M2']['worstA']:.0f} minutes. ", f"Expected time for an Alta X searching every {R2} min (90% per-pass detection, assumed) to find a new spot fire over Camp's standard day, from the sustained to the gust reading of its limit.")
inside = [lab for lab, v in (("SwitchBlade-Elite", NUM["M2"]["sb"]), ("a Type 1 at 30 mph", NUM["M2"]["t30"])) if v <= 85]
miss = [lab for lab, v in (("the Alta X", NUM["M2"]["worstA"]), ("the SwitchBlade-Elite", NUM["M2"]["sb"]), ("a Type 1 at 30 mph", NUM["M2"]["t30"])) if v > 85]
B(d, "Against Camp's 85 minutes. ", f"The first burnovers came 85 min after the first 911 call. Read on gusts, expected detection is {NUM['M2']['worstA']:.0f} min for the Alta X, "
  f"{NUM['M2']['sb']:.0f} for the SwitchBlade-Elite and {NUM['M2']['t30']:.0f} for a Type 1 at 30 mph: " + ("only " + " and ".join(inside) + " detect inside it." if inside else "none detects inside it."))
B(d, "No spot fire record yet. ", "No drone in the library has an operational spot fire detection record (Week 6 Platform, M2 PARTIAL), so per-pass detection runs as 50%, 80% and 95%.")
H(d, "Who uses it")
P(d, ("Lookouts and division supervisors: ", "b"), "a spot across the line changes escape routes and safety zones within minutes. ",
  ("Incident commander and operations: ", "b"), "detection time decides whether initial attack resources reach a spot before it becomes a new front.")
H(d, "Drone types serving M2")
T(d, ["Drone", "Type", "Parameters", "Wind limit", "M2 role"], [
    ["Type 1 and 2 FW (FVR-90, ScanEagle, Aerosonde class)", "NWCG Type 1 / 2", "Type 1: 6 to 14 h, mid-wave IR; Type 2: 1 to 6 h, long-wave IR; above all other incident aircraft", "Not found", "Wide-area IR search, day and night"],
    ["Alta X / SwitchBlade-Elite", "NWCG Type 3", "17 to 18 min average sorties; 5 mi range", "22 / 25 mph", "Initial attack size-up; confirm and geolocate hot spots"],
    ["DJI Mavic 2 Enterprise Advanced (sensor reference)", "Type 4 band, research", "640 x 512 thermal, 8 to 14 µm, plus 4K RGB; 60 m AGL in FLAME 2; thermal -40 to 550 °C low gain", "Not in library", "Fused RGB-IR flame detection benchmark"],
], [1.6, 1.0, 2.55, 0.8, 1.35], size=9)
H(d, "How the uncertainty is quantified")
P(d, ("Time to detect: ", "b"), "a pass every R min that flies with p (wind gate) and sees the spot with Pd finds it on pass G ~ Geometric(q), q = p x Pd: ",
  ("E[T_detect] = R/2 + R (1 - q)/q.", "b"), " Smallest detectable spot ", ("A_min = (k x GSD)^2, GSD = 2h tan(HFOV/2)/N.", "b"),
  f" With an assumed 40° HFOV and k = 3 pixels, a 640-pixel array gives GSD {gsd(60) * 100:.1f} cm at 60 m AGL (A_min {(3 * gsd(60)) ** 2:.2f} m²) and {gsd(1066.8):.2f} m at 3,500 ft (A_min {(3 * gsd(1066.8)) ** 2:.0f} m²).")
res = []
for name, lab in [(AX, "Alta X, 22 mph"), (SB, "SwitchBlade-Elite, 25 mph"), (T30, "Type 1 if 30 mph"), (T40, "Type 1 if 40 mph")]:
    ps, pg = g("Camp", name, gate="p_fly_sustained"), g("Camp", name)
    res.append([lab] + [f"{tdet(R2, ps, Pd):.0f} to {tdet(R2, pg, Pd):.0f}" for Pd in (0.5, 0.8, 0.95)])
T(d, [f"Camp 8 to 9 Nov, R = {R2} min (assumed)", "Pd 50%: min", "Pd 80%: min", "Pd 95%: min"], res, [2.3, 1.45, 1.45, 1.45])
P(d, ("Kinds of uncertainty: ", "b"), "p aleatory (gusts); Pd epistemic until flight-tested in smoke and at night; HFOV and k assumed. ",
  ("v1.2 needs: ", "b"), "Pd and false alarm rate by smoke and night class; search area per fire (Week 6 Q2); the Type 1 wind limit.", after=2)
H(d, "Night and heat sensors")
P(d, "Thermal sees flame and hot ground through smoke and darkness where RGB and lookouts cannot; the Mavic 2 EA range saturates at 550 °C, so it locates flame without measuring intensity. "
  "Smoke detection suits early detection and flame detection larger fires (Keerthinathan et al. 2023, Week 6 M2 Q5). ", night_line("Camp"),
  f" Camp night gusts allowed an Alta X in {pct(g('Camp', AX, 'night'))} of night hours and a Type 1 at 30 mph in {pct(g('Camp', T30, 'night'))}. ", NIGHT_CAP)
H(d, "Hour by hour, standard day: Camp, 8 November 06:00 to 9 November 06:00 PST (Jarbo Gap RAWS, on the fire)")
timeline_table(d, rows)
P(d, (SRC, "i"))
d.save(OUT / f"M2_Spot_Fire_Detection_UQ_v{VER}.docx")

# ============================================================ M3 Firefighter safety
d = new_doc()
header_block(d, "M3", "Firefighter Safety", "Version 1.1. Standard day for every mission and fire: 24 h, 06:00 to 06:00 local, on the peak-run day. GOFER not used.")
rows = tl("M3_Camp")
random.seed(20261020); N = 20000


def geom(q):
    n = 1
    while random.random() > q:
        n += 1
        if n > 10000: break
    return n


def p_neg(T_fire, p_fly, egress, R=15, Pd=0.9, s=0.52, retry=2):
    neg = 0
    for _ in range(N):
        q = p_fly * Pd
        td = random.uniform(0, R) + (R * (geom(q) - 1) if q > 0 else 1e9)
        neg += (T_fire - (td + retry * geom(s) + egress)) < 0
    return neg / N


cases = [("Camp first burnovers, Alta X, sustained", 85, g("Camp", AX, gate="p_fly_sustained")),
         ("Camp first burnovers, Alta X, gusts", 85, g("Camp", AX)),
         ("Camp first burnovers, Type 1 if 30 mph, gusts", 85, g("Camp", T30)),
         ("Carr 26 July, any listed drone, gusts", 62, g("Carr", AX))]
res = {c[0]: [p_neg(c[1], c[2], e) for e in (30, 45, 60)] for c in cases}
NUM["M3"] = {k: v for k, v in res.items()}
H(d, "Bottom line")
B(d, "P(M < 0), the chance the fire reaches the crew before they are out, ", f"for a 45-minute escape at Camp's first burnovers: {pct(res[cases[0][0]][1])} with an Alta X if gusts are ignored, {pct(res[cases[1][0]][1])} if not, {pct(res[cases[2][0]][1])} with a Type 1 rated to 30 mph.")
B(d, "Carr: flyable, still at risk. ", f"Mule Mountain gusts stayed at or below 21 mph all standard day, so every listed drone could fly; P(M < 0) for a 45-minute escape is still {pct(res[cases[3][0]][1])}, half from a missed pass, half from relay and egress.")
B(d, "Escape time is the missing measurement. ", "Crew positions and escape routes at the burnovers are not in project records, so results are shown for 30, 45 and 60 minutes.")
H(d, "Who uses it")
P(d, ("Crew level, lookouts and crew bosses: ", "b"), "their number; a warning must arrive with more time than the escape takes. ",
  ("Division supervisors and the safety officer: ", "b"), "trigger points, escape routes and safety zones. ", ("Incident commander: ", "b"), "whether a crew is committed in the conditions.")
H(d, "Drone types serving M3")
T(d, ["Drone", "Type", "Parameters", "Wind limit", "M3 role"], [
    ["FVR-90", "NWCG Type 1 FW", "12 to 18 h; thermal video; flew as airborne relay in TCL 1", "Not found", "Overwatch of crews and escape routes; relay"],
    ["SuperVolo", "VTOL FW research", "8 h; relay at 1,200 to 1,500 ft MSL", "Not found", "Alternate relay behind terrain"],
    ["Alta X", "NWCG Type 3 rotor", "17 min average sortie; under 2,000 ft AGL", "22 mph (secondary)", "Check escape routes, safety zones, landing zones"],
    ["DJI Matrice 200 v2 (research)", "Type 3 / 4 band", "Sonic anemometer profiles near the fire front", "Not in library", "Local wind for trigger points"],
], [1.35, 1.1, 2.4, 0.95, 1.5], size=9)
H(d, "How the uncertainty is quantified")
P(d, ("M = T_fire - (T_detect + T_notify + T_egress),", "b"), " headline P(M < 0). T_detect: pass every 15 min (assumed), flies with wind gate p, sees the hazard with Pd 0.9 (assumed). "
  "T_notify: retries every 2 min (assumed), each delivered with s = 52%, the TCL 1 telemetry sharing rate with four ground nodes (Wu et al. 2025, pp.9-11), a stand-in for alert delivery. "
  "T_fire: 85 min at Camp (06:25 call to 07:50 burnovers), 62 min at Carr (18:38 aerial warning to 19:40 mayday). Monte Carlo, 20,000 draws per cell.")
T(d, ["P(M < 0) by escape time", "30 min", "45 min", "60 min"], [[c[0]] + [pct(x) for x in res[c[0]]] for c in cases], [3.3, 1.1, 1.1, 1.1])
P(d, ("Kinds of uncertainty: ", "b"), "p aleatory (gusts); Pd, pass interval and retry assumed; s measured once in TCL 1; T_egress epistemic and the largest. ",
  ("v1.2 needs: ", "b"), "crew positions and escape routes at the Camp and Carr burnovers; radio and dispatch logs for real T_notify; a measured alert delivery rate.", after=2)
H(d, "Night")
P(d, night_line("Camp"), f" Camp's burnovers came in daylight, but night wind still grounded an Alta X in {pct(1 - g('Camp', AX, 'night'))} of Camp night hours by gusts. "
  "Crews working at night lose the lookout's view, so a thermal drone overhead is the only continuous watch of the edge. ", NIGHT_CAP)
H(d, "Hour by hour, standard day: Camp, 8 November 06:00 to 9 November 06:00 PST (Jarbo Gap RAWS)")
timeline_table(d, rows)
P(d, (SRC, "i"))
d.save(OUT / f"M3_Firefighter_Safety_UQ_v{VER}.docx")

# ============================================================ M4 Water drop suppression
d = new_doc()
header_block(d, "M4", "Water Drop Suppression", "Version 1.1. Standard day for every mission and fire: 24 h, 06:00 to 06:00 local, on the peak-run day. GOFER not used.")
rows = tl("M4_LNU")
HREL = 30.0; tfall = math.sqrt(2 * HREL / 9.81)


def drift(mph):
    return mph / 2.236936 * tfall


HAWK = 34.5
HOURLY = list(csv.DictReader(open(HERE / "weather_eval" / "critical_window_hourly.csv")))


def hawk(f):
    hr = [r for r in HOURLY if r["fire"] == f and r["station"] == REP[f] and r["gust"]]
    return sum(float(r["gust"]) <= HAWK for r in hr) / len(hr)


camp = WX["Camp"]["JBGC1"]; lnu = WX["LNU"]["ATLC1"]
nights = [r for r in rows if r["period"] == "Night" and r["gust"]]
NUM["M4"] = dict(drift_lo=drift(camp["crit_sust_max_mph"]), drift_hi=drift(camp["crit_gust_max_mph"]), lnu_night_hawk=sum(float(r["gust"]) <= HAWK for r in nights) / len(nights), lnu_night_ax=sum(r["alta_x"] == "fly" for r in nights) / len(nights))
H(d, "Bottom line")
B(d, "No drone in the project library has dropped water. ", "Every M4 aircraft is a test or demonstrator (W5), so v1.1 quantifies what the weather allowed, not what was done.")
B(d, f"{NUM['M4']['drift_lo']:.0f} to {NUM['M4']['drift_hi']:.0f} m of drift. ", f"A release from {HREL:.0f} m (assumed) in Camp's standard-day wind ({camp['crit_sust_max_mph']:.0f} mph sustained, {camp['crit_gust_max_mph']:.0f} mph gusts) moves the load that far before any aiming correction.")
B(d, "LNU's overnight run was flyable for a 30 kt aircraft. ", f"From 20:00 to 06:00 on 18 to 19 August, Atlas Peak gusts allowed a 30 kt (34.5 mph) aircraft in {pct(NUM['M4']['lnu_night_hawk'])} of hours and an Alta X in {pct(NUM['M4']['lnu_night_ax'])}, while the first complex ICS-209 did not exist until 08:00 on 19 August.")
H(d, "Who uses it")
P(d, ("Incident commander, operations section chief and air tactical group supervisor: ", "b"), "they order and sequence drops and need the chance a drop lands on target, especially when crewed aircraft stand down at night. ",
  ("Crews on the line: ", "b"), "they work under the drop, so drift and timing are also a safety number.")
H(d, "Aircraft for M4 (W5 unless noted)")
T(d, ["Aircraft", "Status", "Load", "Wind / night", "M4 role"], [
    ["Rain and Sikorsky autonomous Black Hawk / Firehawk", "Demonstrator; CAL FIRE agreement 2025", "324 gal bucket (demo); 1,000 gal tank class", "Gusts to 30 kt; aimed at night and smoke", "Night direct attack at helicopter scale"],
    ["Seneca Argo-1 swarm (5 drones)", "Tested by CAL FIRE, July 2026", "20 gal water plus foam each", "Not found; IR targeting", "Initial attack on small fires"],
    ["FireSwarm Thunder Wasp GT", "Tested (XPRIZE finalist)", "Up to 100 gal", "Not found", "Initial attack; bulk water"],
    ["Large UAS or optionally piloted aircraft", "Concept (ConOps 2026, p.19)", "Bulk water", "Night by design", "Night water drop instead of retardant"],
], [1.6, 1.25, 1.35, 1.6, 1.25], size=9)
H(d, "How the uncertainty is quantified")
P(d, ("Water on target per hour: ", "b"), ("V_eff = n x S x V x P(hit) x p,", "b"), " n aircraft, S = 60 / (2 transit + drop + turnaround) sorties per hour, V load, p wind gate. "
  f"Drift d = u x sqrt(2h/g) = u x {tfall:.2f} s at h = {HREL:.0f} m; P(hit) = Φ((w - d_res)/σ) - Φ((-w - d_res)/σ) for target half-width w, aim error σ and residual drift d_res. Wind enters through p and d.")
res = []
for f in ["Camp", "LNU", "Carr", "Pioneer"]:
    r = WX[f][REP[f]]
    res.append([f, f"{r['crit_sust_max_mph']:.0f} / {r['crit_gust_max_mph']:.0f}", f"{drift(r['crit_sust_max_mph']):.0f} to {drift(r['crit_gust_max_mph']):.0f}",
                f"{drift(r['crit_gust_max_mph']) - drift(r['crit_sust_max_mph']):.0f}", pct(hawk(f)), pct(g(f, AX, "night"))])
T(d, ["Critical window", "Sust. / gust mph", "Drift m (30 m)", "Gust spread m", "p, 30 kt", "p Alta X, nights"], res, [1.2, 1.2, 1.2, 1.2, 1.0, 1.2])
P(d, ("Kinds of uncertainty: ", "b"), "the gust spread is drift an aim on sustained wind cannot remove; this is a no-drag, instant-wind estimate from 20 ft winds, and drag and stronger flight-level wind push it in opposite directions. Throughput: Week 6 records about 35 L/h for a small drone against 10,345 L/h for a large airtanker. ",
  ("v1.2 needs: ", "b"), "measured drop accuracy against wind from the Seneca and Firehawk tests; release heights and loads per aircraft.", after=2)
H(d, "Night")
P(d, "In the ConOps future state a night water drop would only require a large UAS or optionally piloted aircraft, and the same water could be shared across many small to medium UAS hitting the flame front in a daisy chain (ConOps 2026, p.19). "
  "No aircraft in project records has done it. ", night_line("LNU"), " ", night_line("Camp"))
H(d, "Hour by hour, standard day: LNU, 18 August 06:00 to 19 August 06:00 PDT (Atlas Peak RAWS)")
timeline_table(d, rows)
P(d, (SRC, "i"))
d.save(OUT / f"M4_Water_Drop_Suppression_UQ_v{VER}.docx")

# ============================================================ M5 Logistics support
d = new_doc()
header_block(d, "M5", "Logistics Support", "Version 1.1. Standard day for every mission and fire: 24 h, 06:00 to 06:00 local, on the peak-run day. GOFER not used.")
rows = tl("M5_Pioneer")
V = 50.0; END, RES = 22.0, 5.0; HOLD, TURN = 2.0, 10.0; DIST = 3.0


def radius(w):
    return 0.0 if w >= V else ((END - RES) / 60) / (1 / (V - w) + 1 / (V + w))


def cycle(dk, w):
    return float("inf") if w >= V else 60 * (dk / (V - w) + dk / (V + w)) + HOLD + TURN


def p_on(f, field, kind="crit"):
    if kind == "crit":
        hr = [r for r in HOURLY if r["fire"] == f and r["station"] == REP[f] and r[field]]
    ok = [float(r[field]) <= 22 and cycle(DIST, float(r[field]) * 1.609) <= 30 for r in hr]
    return sum(ok) / len(ok)


pon = {f: (p_on(f, "sust"), p_on(f, "gust")) for f in ["Camp", "LNU", "Pioneer", "Carr"]}
NUM["M5"] = dict(r0=radius(0), r_camp=radius(camp["crit_sust_max_mph"] * 1.609), camp_sust=camp["crit_sust_max_mph"], pon=pon, per_h=60 / cycle(DIST, 0))
H(d, "Bottom line")
B(d, f"{radius(0):.1f} km becomes {radius(camp['crit_sust_max_mph'] * 1.609):.1f} km, then zero. ", f"Loaded radius of an Alta X with 22 lb (10 kg, about 10 L of water): still air, Camp's {camp['crit_sust_max_mph']:.0f} mph sustained headwind, and its {camp['crit_gust_max_mph']:.0f} mph gusts (above the assumed {V:.0f} km/h cruise; its 22 mph limit grounds it first).")
B(d, "P(on time), share of standard-day hours: ", f"Camp {pct(pon['Camp'][1])} to {pct(pon['Camp'][0])}, LNU {pct(pon['LNU'][1])} to {pct(pon['LNU'][0])}, Pioneer {pct(pon['Pioneer'][1])} to {pct(pon['Pioneer'][0])}, for a 3 km delivery inside 30 minutes (gust to sustained reading).")
B(d, "Small drones carry litres, crews need loads. ", f"About {60 / cycle(DIST, 0):.1f} deliveries per hour, {10 * 60 / cycle(DIST, 0):.0f} kg/h per Alta X; heavy-lift aircraft in development (Elroy Chaparral, 500 lb over 300 nm; Piasecki KARGO II, 1,500 lb) are the M5 class.")
H(d, "Who uses it")
P(d, ("Crew bosses and division supervisors: ", "b"), "whether water, hose, pumps and food arrive when needed on remote lines (Pioneer) or when roads are jammed by evacuation (Camp). ",
  ("Logistics section chief and incident commander: ", "b"), "how many aircraft to order, which crews to supply by air, and which helicopter landing zones stay usable.")
H(d, "Aircraft serving M5")
T(d, ["Aircraft", "Status", "Payload and endurance", "Wind limit", "M5 role"], [
    ["Alta X (Freefly)", "In service, flown with ACERO", "35 lb max; 50 min empty, 35 min at 8 lb, 22 min at 22 lb; 60 mph top speed", "22 mph (secondary)", "Light resupply under 10 kg; landing zone checks"],
    ["Type 3 / 4 agency small UAS", "In service", "Light loads; Type 3 range 5 mi, Type 4 under 2 mi", "Not found", "Radios, batteries, first aid"],
    ["Elroy Air Chaparral", "In development (W5)", "500 lb over 300 nm", "Not found", "Crew resupply from staging"],
    ["Piasecki KARGO II", "In development (W5)", "1,500 lb", "Not found", "Water, hose and pumps in bulk"],
], [1.4, 1.3, 2.35, 0.95, 1.3], size=9)
H(d, "How the uncertainty is quantified")
P(d, ("P(on time) = p x P(complete | fly) x P(cycle <= window).", "b"), " Loaded radius into headwind w at airspeed V: r = t_usable / (1/(V - w) + 1/(V + w)). "
  "Cycle = round trip + hold + turnaround (Week 6 M5 Q4). Endurance 22 min at 22 lb (W5); reserve 5 min, V = 50 km/h, hold 2 min, turnaround 10 min, crew at 3 km all assumed. "
  "P(complete | fly) = 1 until a delivery success rate is measured. P(on time) is the share of critical-window hours meeting both the wind limit and the window.")
res = []
for f in ["Camp", "LNU", "Pioneer", "Carr"]:
    r = WX[f][REP[f]]; ws, wg = r["crit_sust_max_mph"] * 1.609, r["crit_gust_max_mph"] * 1.609
    cs, cg = cycle(DIST, ws), cycle(DIST, wg)
    res.append([f, f"{radius(ws):.1f} / {radius(wg):.1f}", f"{cs:.0f} / {'none' if math.isinf(cg) else f'{cg:.0f}'}", f"{pct(pon[f][1])} to {pct(pon[f][0])}", pct(g(f, AX, "night"))])
T(d, ["Alta X, 10 kg", "Radius km at peak sust. / gust", "Cycle min at peak sust. / gust", "P(on time), standard-day hours", "p gust, nights"], res, [1.3, 1.6, 1.6, 1.45, 1.2])
P(d, ("Kinds of uncertainty: ", "b"), "p and w aleatory (station wind); airspeed, turnaround and success rate assumed; crew demand is not in project records, so fleet size is not computed. ",
  ("v1.2 needs: ", "b"), "crew demand per shift; measured Alta X loaded cruise speed and wind limit; distances from staging to crews on Pioneer and Camp.", after=2)
H(d, "Night")
P(d, night_line("Pioneer"), f" Pioneer night gusts allowed an Alta X in {pct(g('Pioneer', AX, 'night'))} of night hours and LNU's in {pct(g('LNU', AX, 'night'))}, "
  "so wind is rarely what stops a night delivery; authorisation is. ", NIGHT_CAP)
H(d, "Hour by hour, standard day: Pioneer, 29 August 06:00 to 30 August 06:00 MDT (Bearskin Creek RAWS)")
timeline_table(d, rows)
P(d, (SRC, "i"))
d.save(OUT / f"M5_Logistics_Support_UQ_v{VER}.docx")

json.dump(NUM, open(OUT / f"mission_uq_v{VER}_numbers.json", "w"), indent=1, default=str)
print("wrote", sorted(p.name for p in OUT.iterdir()))
