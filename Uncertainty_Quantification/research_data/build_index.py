"""Builds Research_Data_Index.xlsx from entrapments/ and datasets/.
Run: uv run --project ~/nasa_rag python build_index.py"""
import json, csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
H = Path(__file__).parent
E, D = H / "entrapments", H / "datasets"

def sheet(ws, header, rows, widths):
    ws.append(header); [ws.append(r) for r in rows]
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
    for row in ws.iter_rows():
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"

wb = Workbook()
# 1. Curated timing anchors for the safety margin M = T_fire - (T_detect + T_notify + T_egress)
sheet(wb.active, ["Fire", "Event", "Clock time", "What happened", "Margin term", "Source", "Checked against text"], [
 ["Camp 2018", "Rattlesnake Flats burnover", "~14:15", "Fire estimated 1/4 mile north of the road, backing south; firing plan organised", "T_fire (initial distance)", "[CAL FIRE 2018, p.8]", "yes"],
 ["Camp 2018", "Rattlesnake Flats burnover", "~14:45", "Wind suddenly shifted, ~25 mph, 10 to 15 ft flames; fire crossed road, egress cut off both ways", "T_fire", "[CAL FIRE 2018, p.8]; [CAL FIRE 2018, p.9]", "yes"],
 ["Camp 2018", "Rattlesnake Flats burnover", "15:15, 15 min", "NIST logs the burnover at 15:15, duration 15 min: 30 min later than the Green Sheet wind shift (source timing uncertainty)", "timing error prior", "[Maranghides et al 2021, p.197]", "yes"],
 ["Camp 2018", "Rattlesnake Flats burnover", "~15:00", "Emergency traffic call went unanswered due to radio traffic volume", "T_notify (failure)", "[CAL FIRE 2018, p.9]", "yes"],
 ["Camp 2018", "Rattlesnake Flats burnover", "", "61 F, RH 12%, NE 10 mph gusting 28, visibility good; 3 seriously burned", "weather input", "[CAL FIRE 2018, p.7]", "yes"],
 ["Camp 2018", "Fire behaviour, first 12 h", "06:29 onward", "~15 miles downwind, >55,000 ac, ~100 ch/h; spotting up to 1 mile ahead", "T_fire (ROS)", "[CAL FIRE 2018, p.7]", "yes"],
 ["Camp 2018", "Hoffman Rd burnover", "07:35 to 07:50", "Large spot reported 07:35; within 10 min road blocked, firefighters and civilians trapped, shelters used as shields", "T_detect to T_fire gap", "[Maranghides et al 2021, p.198]", "yes"],
 ["Camp 2018", "19 burnovers", "Nov 8 07:50 to Nov 9 07:15", "Durations 15 to 360 min; see Camp Burnovers tab", "exposure duration", "[Maranghides et al 2021, p.197]", "yes"],
 ["Carr 2018", "Wind-driven transition", "~17:30", "Fire moved from fuel-dominated to wind-driven", "T_fire", "Carr_2018_CALFIRE_GreenSheet (entrapments/)", "yes"],
 ["Carr 2018", "Dozer burnover", "~18:10 / 18:13", "Burnover reported, then relayed by DIVS", "T_fire", "Carr_2018_CALFIRE_GreenSheet", "yes"],
 ["Carr 2018", "Fire whirl, Branch III", "18:38", "Air Tactical Group Supervisor advised Branch III to evacuate", "T_detect / T_notify (aerial warning)", "Carr_2018_CALFIRE_GreenSheet", "yes"],
 ["Carr 2018", "Fire whirl, Branch III", "19:23", "Fire front crested the hill and advanced", "T_fire", "Carr_2018_CALFIRE_GreenSheet", "yes"],
 ["Carr 2018", "Fire whirl, Branch III", "19:40", "Fire Prevention Inspector radioed mayday (62 min after the aerial warning)", "margin outcome", "Carr_2018_CALFIRE_GreenSheet", "yes"],
 ["Carr 2018", "Weather", "18:00", "RAWS 111 F", "weather input", "Carr_2018_CALFIRE_GreenSheet", "yes"],
 ["Carr 2018", "Source conflict", "18:40", "Blue Sheet gives ~6:40 pm for the dramatic increase in fire behaviour", "timing error prior", "Carr_2018_CALFIRE_BlueSheet", "yes"],
 ["Cedar 2003", "Engine 6162, Wynola", "Oct 29 afternoon", "Run of 0.46 mi upslope in a little under 2 minutes; wind 17 mph gusting 31 from the west", "T_fire (extreme ROS)", "Cedar_2003_CDF_GreenSheet; NIOSH report", "yes"],
 ["Cedar 2003", "Engine 6162, Wynola", "", "About 20 s before conditions changed dramatically; ~170 ft to the refuge (house rear door)", "T_egress vs T_fire", "Cedar_2003_NovatoFPD report; Cedar_2003_CDF_Report", "yes"],
 ["Eaton 2025", "West Altadena evacuation (civilian)", "02:18 / 03:25 / ~05:13", "Evacuation tasks initiated 02:18, orders released 03:25, conflagration ~05:13", "T_notify (warning pipeline)", "Eaton_2026_LACounty_InvestigationReport", "yes"],
 ["Pioneer 2016", "No entrapment found", "", "Only a fuel geyser burn (medical) and a radio cloning close call (comms)", "T_notify (comms failure mode)", "Pioneer_2016_* files", "per agent, one day's search"],
 ["LNU 2020", "No entrapment found", "", "No firefighter entrapment report located", "", "entrapment_sources.json", "one day's search"],
], [12, 28, 16, 60, 22, 40, 14])
wb.active.title = "Margin Timing Anchors"

rows = list(csv.reader(open(E / "Camp_2018_NIST_burnovers.csv")))
sheet(wb.create_sheet("Camp Burnovers"), rows[0], rows[1:], [5, 30, 26, 10, 40, 10, 11, 10, 30, 70, 26])

es = json.load(open(E / "entrapment_sources.json"))
k = list(es[0].keys())
sheet(wb.create_sheet("Entrapment Sources"), k, [[str(x.get(c, "")) for c in k] for x in es], [14, 6, 30, 12, 20, 20, 24, 40, 30, 30, 30, 40, 40, 30, 11, 8])

dc = json.load(open(D / "dataset_catalog.json"))
k2 = list(dc[0].keys())
sheet(wb.create_sheet("Dataset Catalog"), k2, [[json.dumps(x.get(c)) if isinstance(x.get(c), (list, dict)) else str(x.get(c, "")) for c in k2] for x in dc], [18, 30, 20, 20, 16, 16, 12, 40, 12, 30, 9, 30, 40, 8, 11][:len(k2)] + [20] * max(0, len(k2) - 15))

sheet(wb.create_sheet("Gaps"), ["Gap", "Effect on UQ", "Next step"], [
 ["Crew egress times not in library; no public raw firefighter GPS tracks", "T_egress rests on published travel-rate models plus elicitation", "Use Sullivan 2020 and Campbell 2017/2019 coefficients; add egress questions to questionnaire"],
 ["Page et al 2019 entrapment database link returns 404", "No national point dataset for P(entrapment) priors", "Email authors / RMRS for the file"],
 ["No firefighter entrapment on Pioneer, LNU, Eaton (one day's search, 2026-09-30)", "Camp, Carr, Cedar carry validation of the margin", "Ask NASA team whether Carr and Cedar can join the validation set"],
 ["NIROPS IR perimeters behind FAMAuth since 2026-05-04", "Night perimeter latency cannot be measured from public data", "Request access through the NASA/agency partner"],
 ["RAWS hourly archive needs a WRCC access code; Synoptic needs token", "Ridge-top wind input uncertain; airport stations sit in valleys", "Email wrcc@dri.edu; free Synoptic token"],
 ["LANDFIRE fuels 3+ GB, LF2014 for Pioneer not directly linked", "Fuel map uncertainty (largest ROS bias driver) not yet sampled", "Clip-and-ship per fire when simulations start"],
 ["Eaton overnight run has no public perimeter", "Eaton cannot validate early spread", "Treat Eaton as warning-pipeline case only"],
 ["Link and Maranghides 2023 (23 Camp burnovers, 157 responder accounts) full text not retrieved", "Could extend the Camp burnover table", "Download manually, DOI 10.1071/WF22115"],
], [55, 50, 55])
wb.save(H / "Research_Data_Index.xlsx"); print("ok")
