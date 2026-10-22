"""Two slides per mission document (M1 to M5, version 1.1), in the v0.5 deck's look.

Slide A: the bottom line as stat cards (the v0.5 slide 2 format), who uses it, which drones.
Slide B: the hour-by-hour strip from the perimeter records and station wind (fly or not by gust for the Alta X
and a Type 1 at 30 mph, newest aircraft perimeter age, drone task phases), key events, and the UQ result.
All numbers come from build_mission_uq_v11.py outputs and timeline_eval/.

Run after build_mission_uq_v11.py:
    uv run --with python-pptx python build_mission_uq_v11_deck.py
"""
import csv, json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = Path(__file__).parent
OUT = HERE / "Mission UQ v1"
NUM = json.load(open(OUT / "mission_uq_v1.1_numbers.json"))
WX = json.load(open(HERE / "weather_eval" / "results.json"))
NIGHT = json.load(open(HERE / "timeline_eval" / "night_stats.json"))

C = {k: RGBColor.from_string(v) for k, v in dict(
    ink="1D2125", ink2="3A424B", bg="FFFFFF", bg2="F1F2F4", ember="E8642F", blue="2A78D6", green="18A878",
    violet="8E7BD0", amber="E9A821", slate="6B7480", red="C8402A", grid="E1E4E8", night="2F3A52").items()}
FOOT = "NASA Aerial Wildfire Project  |  ACERO Disaster Lab  |  Mission UQ v1.1  |  20 October 2026"
W, M = 13.333, 0.6

prs = Presentation(); prs.slide_width = Inches(W); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def box(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0] = 0.06
    return sh


def text(s, x, y, w, h, runs, size=14, color=None, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Calibri"):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"): setattr(tf, side, 0)
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
        p.space_after = Pt(4)
        for seg in (para if isinstance(para, list) else [para]):
            t, o = (seg, {}) if isinstance(seg, str) else seg
            r = p.add_run(); r.text = t
            r.font.size = Pt(o.get("size", size)); r.font.bold = o.get("bold", bold); r.font.name = o.get("font", font)
            r.font.color.rgb = o.get("color", color or C["ink"])
    return tb


def frame(kicker, title, n):
    s = prs.slides.add_slide(BLANK)
    text(s, M, 0.35, 9, 0.3, kicker.upper(), size=12, bold=True, color=C["ember"], font="Arial")
    text(s, M, 0.62, 12.1, 0.75, title, size=28, bold=True, font="Arial")
    text(s, M, 7.05, 10, 0.28, FOOT, size=10, color=C["slate"])
    text(s, 12.2, 7.05, 0.55, 0.28, str(n), size=10, color=C["slate"], align=PP_ALIGN.RIGHT)
    return s


def stat(s, x, y, w, big, label, color):
    box(s, x, y, w, 2.35, fill=C["bg2"], shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + 0.25, y + 0.15, w - 0.5, 0.85, big, size=34, bold=True, color=color, font="Arial", anchor=MSO_ANCHOR.BOTTOM)
    text(s, x + 0.25, y + 1.08, w - 0.5, 1.2, label, size=13, color=C["ink2"])


def slide_a(n, mid, title, kicker, stats, who, drones):
    s = frame(f"{mid}  |  {kicker}  |  bottom line", title, n)
    w, gap = 3.9, 0.2
    for i, (big, label, col) in enumerate(stats):
        stat(s, M + i * (w + gap), 1.6, w, big, label, col)
    box(s, M, 4.2, 5.9, 2.6, fill=C["bg2"], shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, M + 0.25, 4.32, 5.4, 0.35, "Who uses it", size=15, bold=True, font="Arial")
    text(s, M + 0.25, 4.75, 5.4, 2.0, [[(a, {"bold": True}), b] for a, b in who], size=13)
    box(s, M + 6.1, 4.2, 6.0, 2.6, fill=C["bg2"], shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, M + 6.35, 4.32, 5.5, 0.35, "Drone types and parameters", size=15, bold=True, font="Arial")
    text(s, M + 6.35, 4.75, 5.5, 2.0, [[(a, {"bold": True}), b] for a, b in drones], size=12)
    return s


def slide_b(n, mid, kicker, title, csvname, events, result, note):
    s = frame(f"{mid}  |  {kicker}  |  hour by hour", title, n)
    rows = list(csv.DictReader(open(HERE / "timeline_eval" / f"{csvname}_hourly.csv")))
    nh = len(rows); x0, lab = M + 1.55, 1.5
    cw = (W - M - x0) / nh
    bands = [("Hour", 0.32), ("Gust mph", 0.36), ("Alta X 22", 0.32), ("Type 1 at 30", 0.32), ("Map age h", 0.36)]
    y = 1.55
    for name, h in bands:
        text(s, M, y, lab, h, name, size=11, bold=True, color=C["ink2"], anchor=MSO_ANCHOR.MIDDLE)
        for i, r in enumerate(rows):
            x = x0 + i * cw
            night = r["period"] == "Night"
            if name == "Hour":
                box(s, x, y, cw - 0.02, h, fill=C["night"] if night else C["bg2"])
                text(s, x, y, cw - 0.02, h, r["hour"][-5:-3], size=10, bold=True, color=C["bg"] if night else C["ink"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            elif name == "Gust mph":
                v = r["gust"] or "n/a"
                hot = r["gust"] and float(r["gust"]) > 30
                box(s, x, y, cw - 0.02, h, fill=C["bg2"])
                text(s, x, y, cw - 0.02, h, v, size=10, bold=bool(hot), color=C["red"] if hot else C["ink"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            elif name in ("Alta X 22", "Type 1 at 30"):
                v = r["alta_x"] if name == "Alta X 22" else r["type1_30"]
                col = C["green"] if v == "fly" else (C["red"] if v == "no" else C["grid"])
                box(s, x, y, cw - 0.02, h, fill=col)
            else:
                a = r["aircraft_perimeter"]
                v = "none" if a == "none yet" else a.split(",")[-1].replace(" h old", "").strip()
                text(s, x, y, cw - 0.02, h, v, size=10, color=C["ink"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        y += h + 0.06
    # drone task phases as a band of labelled spans
    text(s, M, y, lab, 0.62, "Drone task", size=11, bold=True, color=C["ink2"], anchor=MSO_ANCHOR.MIDDLE)
    spans, start = [], 0
    for i in range(1, nh + 1):
        if i == nh or rows[i]["task"] != rows[start]["task"]:
            spans.append((start, i, rows[start]["task"])); start = i
    cols = [C["blue"], C["ember"], C["violet"], C["green"]]
    for k, (a, b, t) in enumerate(spans):
        box(s, x0 + a * cw, y, (b - a) * cw - 0.03, 0.62, fill=cols[k % 4], shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        span_w = (b - a) * cw
        if span_w < 1.3:  # narrow span: first clause, at most four words
            import re
            t = " ".join(re.split(r"[:;,]", t)[0].split()[:4])
        text(s, x0 + a * cw + 0.05, y + 0.02, span_w - 0.12, 0.58, t, size=8 if span_w < 1.2 else (9 if span_w < 2.5 else 10), color=C["bg"], anchor=MSO_ANCHOR.MIDDLE)
    y += 0.75
    text(s, M, y, 12.1, 0.25, "Green: can fly by station gusts. Red: grounded. Dark hour cells: night (20:00 to 06:00). Map age: hours since the newest aircraft perimeter in the record.",
         size=10, color=C["slate"])
    y += 0.4
    box(s, M, y, 7.3, 6.85 - y, fill=C["bg2"], shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, M + 0.2, y + 0.1, 6.9, 0.3, "What the records show", size=14, bold=True, font="Arial")
    text(s, M + 0.2, y + 0.45, 6.9, 6.85 - y - 0.5, [[(a + "  ", {"bold": True}), b] for a, b in events], size=12)
    box(s, M + 7.5, y, 4.6, 6.85 - y, fill=C["ink"], shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, M + 7.7, y + 0.1, 4.2, 0.3, "Quantified uncertainty", size=14, bold=True, color=C["amber"], font="Arial")
    text(s, M + 7.7, y + 0.45, 4.2, 6.85 - y - 0.5, [result, (note, {"size": 11, "color": C["grid"]})], size=13, color=C["bg"])
    return s


def pct(x):
    return "under 1%" if 0 < x < 0.005 else f"{round(100 * x)}%"


# ---------------------------------------------------------------- title
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, W, 7.5, fill=C["ink"])
text(s, M, 2.0, 11.6, 1.9, "Five Missions, Quantified", size=40, bold=True, color=C["bg"], font="Arial", anchor=MSO_ANCHOR.BOTTOM)
text(s, M, 4.1, 11.2, 1.0, "Mission UQ v1.1: how each mission's uncertainty is calculated, which drones fit, the weather on the fires, and where drones fit hour by hour",
     size=20, color=C["bg2"])
text(s, M, 6.6, 11, 0.4, "NASA Aerial Wildfire Project  |  ACERO Disaster Lab working group  |  20 October 2026  |  Standard day: 24 h, 06:00 to 06:00, peak-run day  |  GOFER not used", size=13, color=C["slate"])

n = 2
m1 = NUM["M1"]
slide_a(n, "M1", "The map is the uncertainty: how old it is, and whether drones can keep it fresh", "Fireline mapping", [
    (f"{m1['a_air']:.0f} h", f"Oldest the newest aircraft perimeter got during Carr's 26 July run; every listed drone could fly all 24 hours", C["ember"]),
    (f"{pct(m1['pA'])} of hours", "Camp standard day: hours an Alta X could fly by gusts at Jarbo Gap, so its map would be most of the day old", C["red"]),
    (f"{m1['age30']:.1f} h", f"Expected map age at Camp for a Type 1 rated to 30 mph on a 1 h revisit (about {round(m1['unseen30'], -2):,.0f} ac unmapped)", C["green"])],
    [("Incident commander and planning: ", "map age decides whether resource orders and evacuation triggers rest on a current line."),
     ("Division supervisors and lookouts: ", "head and flank position, refreshed within the hour.")],
    [("FVR-90, Type 1 FW: ", "12 to 18 h, thermal video, wind limit not found"), ("SuperVolo: ", "8 h, relay and second mapping tier"),
     ("Alta X, Type 3: ", "17 min sorties, 22 mph"), ("SwitchBlade-Elite: ", "18 min sorties, 25 mph")]); n += 1
slide_b(n, "M1", "Fireline mapping", "Carr, 26 to 27 July: the line moved while the record stood still", "M1_Carr", [
    ("17:30", "Run east at 2.5 mph, spotting over 1 mile; 17:34 aircraft perimeter 44,484 ac"),
    ("19:17", "Spots east of the Sacramento River; 19:20 to 20:00 vortex on radar"),
    ("19:45", "ICS-209 28,763 ac, unchanged at 20:15 and 21:00"),
    ("00:00", "Next aircraft perimeter 50,050 ac; spots south of Whiskeytown Lake"),
    ("Night", f"{pct(NIGHT['Carr']['night_stamped_share'])} of Carr aircraft perimeters were flown at night, but about {pct(NIGHT['Carr']['night_growth_share'])} of growth fell at night")],
    f"E[A] = R (2 - p) / (2p). Carr: p = 100%, so a 1 h revisit holds the map near 0.5 h old. Camp: {m1['ageS']:.1f} h if the limit is read on sustained wind, {m1['ageA']:.1f} h on gusts.",
    "Largest unknowns: sustained versus gust reading, and the unpublished Type 1 wind limit."); n += 1

m2 = NUM["M2"]
slide_a(n, "M2", "Spot fires: detection time depends on whether the drone can fly at all", "Spot fire detection", [
    (f"{m2['best']:.0f} to {m2['worstA']:.0f} min", "Expected time for an Alta X to find a spot at Camp, from the sustained to the gust reading (30 min search, Pd 0.9 assumed)", C["ember"]),
    ("85 min", "First 911 call to first burnovers at Camp: the clock every detection competes with", C["red"]),
    (f"{m2['t30']:.0f} min", "Same search by a Type 1 rated to 30 mph, read on gusts", C["green"])],
    [("Lookouts and division supervisors: ", "a spot across the line changes escape routes and safety zones within minutes."),
     ("Incident commander: ", "whether initial attack reaches a spot before it becomes a new front.")],
    [("Type 1 / 2 FW: ", "6 to 14 h mid-wave IR, 1 to 6 h long-wave IR"), ("Alta X / SwitchBlade-Elite: ", "initial attack size-up, confirm hot spots"),
     ("Mavic 2 EA (sensor reference): ", "640 x 512 thermal, 8 to 14 µm, saturates at 550 °C")]); n += 1
slide_b(n, "M2", "Spot fire detection", "Camp, 8 to 9 November: initial attack in wind that grounded small drones", "M2_Camp", [
    ("06:25", "First 911 call; 06:44 first engine reports about 35 mph wind; 06:45 second ignition"),
    ("07:10", "Engine estimate 200 to 300 ac; 07:37 Concow order; 07:50 first burnovers"),
    ("08:01", "Air attack: 'fire is now in Paradise', about 2,000 ac"),
    ("09:54", "First aircraft perimeter: 54,586 ac"),
    ("Night", f"Camp night gusts grounded an Alta X in {pct(1 - WX['Camp']['JBGC1']['perimeter_night_gate']['Alta X (22 mph)']['p_fly_gust'])} of night hours")],
    f"E[T] = R/2 + R (1 - q)/q, q = p x Pd. Read on gusts: Alta X {m2['worstA']:.0f} min, SwitchBlade-Elite {m2['sb']:.0f} min, Type 1 at 30 mph {m2['t30']:.0f} min, against 85.",
    "Pd has no operational record yet; v1.1 runs 50%, 80% and 95%."); n += 1

m3 = NUM["M3"]; k = list(m3.keys())
slide_a(n, "M3", "Crew safety: the chance the fire wins, P(M < 0)", "Firefighter safety", [
    (pct(m3[k[0]][1]), "P(M < 0) at Camp's first burnovers, 45 min escape, Alta X overwatch, limit read on sustained wind", C["green"]),
    (pct(m3[k[1]][1]), "Same case with the limit read on gusts: the drone is grounded most of the morning", C["red"]),
    (pct(m3[k[3]][1]), "Carr 26 July, every drone flyable: risk left from a missed pass, relay retries and egress", C["ember"])],
    [("Crews, lookouts, crew bosses: ", "a warning must arrive with more time than the escape takes."),
     ("Division supervisors, safety officer: ", "trigger points, escape routes, safety zones."), ("Incident commander: ", "commit or pull a crew.")],
    [("FVR-90: ", "overwatch and airborne relay (TCL 1)"), ("SuperVolo: ", "relay behind terrain"),
     ("Alta X: ", "escape routes, safety zones, landing zones"), ("Matrice 200 v2: ", "local wind profiles for trigger points")]); n += 1
slide_b(n, "M3", "Firefighter safety", "Camp, 8 November: every crew hazard came before the record", "M3_Camp", [
    ("07:50", "First burnovers on Hoffman and Concow roads, 85 min after the first call"),
    ("14:15", "Firing operation on Rattlesnake Flats Rd; 14:45 wind shift to about 25 mph"),
    ("15:00", "Emergency traffic call unanswered (radio volume); 15:15 burnover"),
    ("18:30", "First ICS-209: 16,000 ac")],
    "M = T_fire - (T_detect + T_notify + T_egress). P(M < 0) by escape time 30 / 45 / 60 min, Camp Alta X on gusts: "
    + " / ".join(pct(x) for x in m3[k[1]]) + "; Type 1 at 30 mph: " + " / ".join(pct(x) for x in m3[k[2]]) + ".",
    "Escape time is the missing measurement: crew positions and routes are not in the records."); n += 1

m4 = NUM["M4"]
slide_a(n, "M4", "Water drops: no drone has done it, so we quantify what the wind allowed", "Water drop suppression", [
    (f"{m4['drift_lo']:.0f} to {m4['drift_hi']:.0f} m", "Drift of a 30 m release (assumed) in Camp's standard-day wind, sustained to gust, before aim correction", C["ember"]),
    (pct(m4["lnu_night_hawk"]), "LNU overnight run hours a 30 kt aircraft (autonomous Firehawk class) could fly by Atlas Peak gusts", C["green"]),
    ("35 vs 10,345 L/h", "Small drone against large airtanker throughput (Week 6)", C["red"])],
    [("Incident commander, operations, air tactical group supervisor: ", "order and sequence drops, especially when crewed aircraft stand down at night."),
     ("Crews on the line: ", "they work under the drop.")],
    [("Autonomous Firehawk: ", "324 gal bucket (demo), gusts to 30 kt"), ("Seneca Argo-1: ", "5 drones, 20 gal each (tested)"),
     ("FireSwarm Thunder Wasp: ", "up to 100 gal (tested)"), ("Large UAS / OPA: ", "night water drop concept (ConOps p.19)")]); n += 1
slide_b(n, "M4", "Water drop suppression", "LNU, 18 to 19 August: the overnight run nobody could see", "M4_LNU", [
    ("18:00", "Atlas Peak 29 mph sustained, 41 mph gusts"),
    ("Night", "Fires merge in a 12-mile overnight run (agency chronology)"),
    ("00:00", "Satellite 41,602 ac for the main object, up from 7,328 ac at 12:00"),
    ("08:00", "First complex ICS-209 (19 Aug): 33,025 ac"),
    ("Record", "No aircraft perimeter series exists for LNU without GOFER")],
    "V_eff = n x S x V x P(hit) x p. Wind enters twice: through p and through drift d = u x sqrt(2h/g).",
    "The gust spread of the drift is what an aim on sustained wind cannot remove."); n += 1

m5 = NUM["M5"]
slide_a(n, "M5", "Logistics: P(on time) for water and supplies to crews", "Logistics support", [
    (f"{m5['r0']:.1f} to {m5['r_camp']:.1f} km", f"Alta X loaded radius with 10 kg, still air to Camp's peak {m5['camp_sust']:.0f} mph sustained headwind; zero in its gusts", C["ember"]),
    (f"{pct(m5['pon']['Camp'][1])} to {pct(m5['pon']['Camp'][0])}", "Camp standard day P(on time), 3 km delivery inside 30 min (gust to sustained reading)", C["red"]),
    (f"{pct(m5['pon']['Pioneer'][1])} to {pct(m5['pon']['Pioneer'][0])}", "Pioneer standard day: remote lines, light wind, so small drones can resupply", C["green"])],
    [("Crew bosses and division supervisors: ", "water, hose, pumps and food on remote lines."),
     ("Logistics section chief and IC: ", "how many aircraft, which crews by air, which helicopter landing zones stay usable.")],
    [("Alta X: ", "22 min at 22 lb, 60 mph top speed, 22 mph limit"), ("Type 3 / 4: ", "radios, batteries, first aid"),
     ("Elroy Chaparral: ", "500 lb over 300 nm (in development)"), ("Piasecki KARGO II: ", "1,500 lb (in development)")]); n += 1
slide_b(n, "M5", "Logistics support", "Pioneer, 29 August: a flyable day to resupply before the run", "M5_Pioneer", [
    ("08:00", "Newest aircraft perimeter 111,658 ac, already 15 h old"),
    ("29 Aug", "About 29,000 ac of growth in the northeastward run"),
    ("18:16", "Aircraft perimeter 140,778 ac; 19:00 ICS-209 111,604 ac, 58% contained"),
    ("Night", f"{pct(NIGHT['Pioneer']['night_stamped_share'])} of Pioneer perimeters flown at night; about {pct(NIGHT['Pioneer']['night_growth_share'])} of growth at night")],
    "P(on time) = p x P(complete | fly) x P(cycle <= window). Wind rarely stopped delivery on Pioneer; authorisation and payload do.",
    "Crew demand per shift is not in the records, so fleet size is not computed yet.")


# ---------------------------------------------------------------- speaker notes: what each number and variable means
STRIP = ("How to read the hour strip: one column per hour of the standard day (06:00 to 06:00 local on the peak-run day). "
         "Gust mph is the peak wind in that hour at the station named on the document (red when above 30 mph). "
         "Alta X 22 and Type 1 at 30 are fly or no-fly: green when that hour's gust is at or under the drone's wind limit, red when above. "
         "Map age h is hours since the newest aircraft perimeter in the record. The coloured bars are the drone task we propose for each phase.")
GATE = ("p, the wind gate, is the share of hours in the standard day when the station wind is at or under the drone's wind limit. "
        "Read on sustained wind (10-minute mean) it is the optimistic value; read on gusts (peak in the hour) it is the conservative value. "
        "Both are measured 20 ft above ground, so wind at flight altitude is stronger and both overstate availability.")
NOTES = [
 "Title. This deck has two slides per mission document, M1 to M5, version 1.1. Every number comes from data in the project folder: station weather (Weather Data), "
 "the perimeter time-area records (perimeter_supplements, GOFER rows removed), the Week 6 workbook and the v0.5 proposal. Each mission is evaluated on the same standard day: "
 "24 hours, 06:00 to 06:00 local, on that fire's peak-run day (Camp 8 to 9 Nov 2018, Carr 26 to 27 Jul 2018, Pioneer 29 to 30 Aug 2016, LNU 18 to 19 Aug 2020).",
 # M1 A
 f"M1 fireline mapping, bottom line. The quantity is information age: how old the newest fireline map is when someone uses it. "
 f"First card: on Carr the newest aircraft perimeter reached {m1['a_air']:.0f} hours old during the run, while station wind would have let every listed drone fly every hour. "
 f"Second card: on Camp an Alta X read against gusts could fly in only {pct(m1['pA'])} of hours, so its map would be most of a day old. "
 f"Third card: a Type 1 fixed wing rated to 30 mph flies {pct(m1['p30'])} of Camp hours and keeps the expected map age near {m1['age30']:.1f} hours. "
 "The formula is on the next slide's notes. " + GATE,
 # M1 B
 "M1 calculation. Expected map age E[A] = R (2 - p) / (2p). "
 "R is the revisit interval: the time between scheduled mapping passes, 1 hour here. "
 "p is the chance a scheduled pass actually flies, the wind gate. "
 "E[A] is the expected age of the newest good map at a random moment. Why the formula: when each pass succeeds with probability p, the time between good maps is R times a geometric number of tries; "
 "averaging the age over that gap gives R (2 - p) / (2p). Check: if p = 1 every pass flies and E[A] = R/2, half the revisit interval. "
 f"Camp: p = 1 on sustained wind gives {m1['ageS']:.1f} h; p = {pct(m1['pA'])} on gusts gives {m1['ageA']:.1f} h, longer than the day itself, which means no useful map. "
 "Map error E(t) = G x A(t): G is the fire's growth rate in acres per hour (Camp: more than 55,000 acres in the first 12 hours, about 4,600 ac/h, from the CAL FIRE Green Sheet in the Camp record), "
 "and A(t) is the map age, so E is the area burned since the map was made. " + STRIP,
 # M2 A
 f"M2 spot fire detection, bottom line. The quantity is the expected time from a spot fire starting to a drone seeing it. "
 f"First card: an Alta X searching every 30 minutes on Camp's standard day takes {m2['best']:.0f} minutes if its wind limit is read on sustained wind and {m2['worstA']:.0f} minutes on gusts, because it is grounded most hours. "
 f"Second card: 85 minutes is the time from the first 911 call (06:25) to the first burnovers (07:50) at Camp, the clock any detection competes with; it is a comparison, not a deadline. "
 f"Third card: a Type 1 rated to 30 mph doing the same search, read on gusts, takes {m2['t30']:.0f} minutes. " + GATE,
 # M2 B
 "M2 calculation. Expected detection time E[T] = R/2 + R (1 - q) / q, with q = p x Pd. "
 "R is the search revisit interval, 30 minutes (assumed). "
 "p is the wind gate, the chance the pass flies. "
 "Pd is the chance the sensor sees the spot on a pass that flies; no drone has an operational record, so it runs at 50%, 80% and 95%, with 90% on the headline numbers. "
 "q is the chance a single pass both flies and detects. R/2 is the average wait for the first pass; R (1 - q) / q is the average extra wait for the failed passes before the first success. "
 "Smallest detectable spot: A_min = (k x GSD) squared, where GSD = 2 h tan(HFOV / 2) / N is the ground size of one pixel. "
 "h is height above ground, HFOV is the camera's horizontal field of view (40 degrees assumed), N is pixels across (640 for the Mavic 2 EA thermal), "
 "and k is how many pixels across a spot must cover to be confirmed (3 assumed). At 60 m that is a pixel of about 7 cm; at 3,500 ft, the Type 1 floor, about 1.2 m. " + STRIP,
 # M3 A
 f"M3 firefighter safety, bottom line. The quantity is P(M < 0), the probability the fire reaches a crew before the crew has finished escaping. "
 f"First card: Camp's first burnovers, 45-minute escape, Alta X overwatch with its limit read on sustained wind: {pct(m3[k[0]][1])}. "
 f"Second card: the same case read on gusts: {pct(m3[k[1]][1])}, because the drone cannot fly most of that day. "
 f"Third card: Carr on 26 July, where every listed drone could fly: {pct(m3[k[3]][1])}; the remaining risk comes from a missed detection pass, warning relay retries and the escape itself. " + GATE,
 # M3 B
 "M3 calculation. Safety margin M = T_fire - (T_detect + T_notify + T_egress). The crew is safe when M is positive. "
 "T_fire is the time from the hazard becoming visible to fire reaching the crew: 85 minutes at Camp (first 911 call 06:25 to first burnovers 07:50) and 62 minutes at Carr (aerial warning 18:38 to mayday 19:40). "
 "T_detect is the time for the drone to see the hazard: passes every 15 minutes (assumed), each flies with probability p (wind gate) and detects with Pd = 0.9 (assumed). "
 "T_notify is the time to get the warning to the crew: a retry every 2 minutes (assumed), each delivered with s = 52%, the TCL 1 telemetry sharing rate with four ground nodes (Wu et al. 2025), used as a stand-in for alert delivery. "
 "T_egress is the time the crew needs to reach a safety zone; crew positions are not in the records, so it is shown at 30, 45 and 60 minutes. "
 "P(M < 0) is estimated by Monte Carlo: 20,000 random draws of T_detect and T_notify per case, counting how often M comes out negative. " + STRIP,
 # M4 A
 f"M4 water drop, bottom line. No drone in the project library has dropped water, so M4 quantifies what the wind would have allowed. "
 f"First card: drift of {m4['drift_lo']:.0f} to {m4['drift_hi']:.0f} m for a load released from 30 m in Camp's standard-day wind, from the sustained to the gust value, before any aiming correction. "
 f"Second card: on LNU's overnight run, a 30 kt (34.5 mph) aircraft, the autonomous Firehawk class, could fly {pct(m4['lnu_night_hawk'])} of night hours by Atlas Peak gusts. "
 "Third card: throughput from the Week 6 workbook, about 35 litres per hour for a small drone against 10,345 for a large airtanker. " + GATE,
 # M4 B
 "M4 calculation. Water on target per hour V_eff = n x S x V x P(hit) x p. "
 "n is the number of aircraft. S is sorties per hour = 60 / (2 x transit + drop + turnaround), with each term in minutes. V is the load per sortie. p is the wind gate. "
 "P(hit) is the chance the drop lands on target: P(hit) = Phi((w - d_res) / sigma) - Phi((-w - d_res) / sigma), where Phi is the standard normal cumulative distribution, "
 "w is half the target width, sigma is the aiming error, and d_res is the drift left after the pilot aims for the wind. "
 "Drift d = u x sqrt(2h / g): u is wind speed in m/s, h is release height (30 m assumed), g is 9.81 m/s2; sqrt(2h/g) is the fall time, 2.47 s from 30 m. "
 "The gust spread, gust drift minus sustained drift, is the part an aim made on sustained wind cannot remove. This is a no-drag, instant-wind estimate; drag and stronger wind aloft pull in opposite directions. " + STRIP,
 # M5 A
 f"M5 logistics, bottom line. The quantity is P(on time): the chance a requested delivery of water or supplies reaches a crew within its time window. "
 f"First card: loaded radius of an Alta X carrying 10 kg, {m5['r0']:.1f} km in still air and {m5['r_camp']:.1f} km into Camp's peak {m5['camp_sust']:.0f} mph sustained headwind; zero in its gusts, which exceed the assumed cruise speed. "
 f"Second card: Camp standard day, a 3 km delivery inside 30 minutes is on time in {pct(m5['pon']['Camp'][1])} to {pct(m5['pon']['Camp'][0])} of hours (gust to sustained reading). "
 f"Third card: Pioneer's lighter wind gives {pct(m5['pon']['Pioneer'][1])} to {pct(m5['pon']['Pioneer'][0])}. " + GATE,
 # M5 B
 "M5 calculation. P(on time) = p x P(complete | fly) x P(cycle <= window). "
 "p is the wind gate. P(complete | fly) is the chance a delivery that launches succeeds; set to 1 until a success rate is measured. "
 "P(cycle <= window) is the chance the round trip fits inside the delivery window (30 minutes assumed). "
 "Loaded radius r = t_usable / (1/(V - w) + 1/(V + w)): t_usable is loaded endurance minus reserve (22 minutes at 22 lb from Week 5, minus 5 minutes reserve), "
 "V is cruise airspeed (50 km/h assumed), w is the headwind; flying out against the wind and back with it is what the two terms add up. "
 "Cycle = round trip + hold (2 min) + turnaround (10 min), the Week 6 M5 Q4 formula, for a crew 3 km away (assumed). "
 "On the slide, P(on time) is the share of standard-day hours that meet both the wind limit and the window. " + STRIP,
]
for slide_, text in zip(prs.slides, NOTES):
    slide_.notes_slide.notes_text_frame.text = text

prs.save(OUT / "Mission_UQ_v1.1_Slides.pptx")
print("slides:", len(prs.slides))
