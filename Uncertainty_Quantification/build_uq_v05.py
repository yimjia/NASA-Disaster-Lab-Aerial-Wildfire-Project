"""Build UQ proposal v0.5 and workbook v0.5.

v0.5 refocuses the v0.4 proposal on three things: what happened on the four fires (when information
arrived versus when the fire grew), how fast the first detection was reported, and which of those
gaps drones can close, using the drone capabilities in the Week 5 (current and future) and Week 6
(library-documented) Platform tabs. Numbers carried from v0.4 come from case_study_evaluation.py and
perimeter_gap_evaluation.py; new numbers come from first_report_evaluation.py.

Run after first_report_evaluation.py:
    uv run --project ~/nasa_rag python build_uq_v05.py
"""
import csv, shutil
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
DATE = "12 October 2026"
CH = {r["fire"]: r for r in csv.DictReader(open(HERE / "first_report_eval" / "first_report_chain.csv"))}

# ======================================================================== document helpers
d = Document()
st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
for s in d.sections:
    s.left_margin = s.right_margin = Inches(0.9); s.top_margin = s.bottom_margin = Inches(0.8)
def H(t, l=1): d.add_heading(t, level=l)
def P(t, italic=False, bold=False, size=None):
    p = d.add_paragraph(); r = p.add_run(t); r.italic = italic; r.bold = bold
    if size: r.font.size = Pt(size)
    return p
def B(items):
    for t in items: d.add_paragraph(t, style="List Bullet")
def N(items):
    for t in items: d.add_paragraph(t, style="List Number")
def T(header, rws, size=8.5, widths=None):
    t = d.add_table(rows=1, cols=len(header)); t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; c.paragraphs[0].add_run(h).bold = True
    for row in rws:
        cells = t.add_row().cells
        for i, v in enumerate(row): cells[i].text = str(v)
    for row in t.rows:
        for i, c in enumerate(row.cells):
            if widths: c.width = Inches(widths[i])
            for p in c.paragraphs:
                for r in p.runs: r.font.size = Pt(size)
    d.add_paragraph()
def FIGURE(path, caption, width=6.6):
    d.add_picture(str(path), width=Inches(width))
    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    P(caption, italic=True, size=9)

# ======================================================================== title and status
tp = d.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = tp.add_run("Quantifying Uncertainty for Firefighter Safety and Aerial Logistics\nin 24/7 Wildland Fire Aerial Operations"); r.bold = True; r.font.size = Pt(17)
sp = d.add_paragraph(); sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
sp.add_run("What happened on four fires, how fast the first information got out, and where drones close the gap\n").italic = True
sp.add_run(f"NASA Aerial Wildfire Project, ACERO Disaster Lab working group\nDraft v0.5, {DATE}")

P("Status. Working draft for team discussion. Literature is cited as [Document Year, p.N] from the project reference "
  "library (NASA RAG). Fire numbers come from the project's own analysis of public records (ICS-209-PLUS, NICC, NIFC GeoMAC, "
  "NASA FEDS, GOFER, FIRMS) and name the record. Drone capabilities come from the Week 6 Platform tab (drones documented "
  "in the library) and the Week 5 Platform tab (current fleet and future aircraft, from agency and vendor sources); Week 5 "
  "figures are labelled as such and are claims, not tests, unless stated. Radio and dispatch logs are still pending. "
  "Novelty has not been assessed.", italic=True, size=9.5)

H("What changed from v0.4")
B([
 "Refocused and shortened: the document now follows the fires in time (Section 2), then the first report (Section 3), "
 "then the information gaps (Section 4) and the drones that could close them (Section 5). Detailed tables moved to the workbook.",
 "New: drone capabilities, current and future, mapped gap by gap to what the four fires show (Section 5), with the wind "
 "limits that decide whether a drone could have flown on the worst day.",
 "New: Camp 2018 perimeter and size record from ICS-209-PLUS, in the team's supplement format "
 "(perimeter_supplements/Camp_2018_Perimeter_Time_Area_ICS209PLUS.xlsx; Section 2.3).",
 "New: the first-report chain for all four fires: first call, first aircraft estimate, first satellite detection, first "
 "perimeter and first ICS-209 (Section 3).",
 "Removed: content from outside this project, the version history, and repeated method notes.",
])

H("Summary")
P("On all four fires the fire outran the record. The official size on the ICS-209 was usually the previous night's "
  "perimeter, so decision makers read a number about one day old, and on blowup days that number was far too small: at "
  "18:30 on the first day of the Camp Fire the first ICS-209 gave 16,000 acres, while CAL FIRE later put the first 12 hours "
  "at more than 55,000 acres and 15 of the 19 documented burnovers had already happened. The first information about each fire "
  "came fast from people (a 911 call, a lead plane, air attack by radio) but slowly as a record: the first mapped perimeter "
  "arrived 3.5 to 14 hours after ignition and the first ICS-209 5 to 12 hours after (50 hours for the LNU complex). Most growth "
  "happened while the newest perimeter was more than 12 hours old (59% to 91% of final area on Carr, Camp and Pioneer), and "
  "37% of LNU's growth happened at night, when crewed aircraft do not fly. The hourly LNU record shows that a 2 to 3 hour "
  "revisit around the clock would keep the last-known size within about 3 to 5% during growth, against 20 to 32% at the 12 to "
  "24 hour cadence of the records. Long-endurance (NWCG Type 1) drones already flown with NASA ACERO can in principle meet "
  "that cadence day and night; small agency drones, whose average sortie is about 17 minutes, cannot. The binding limit is "
  "wind: on the first morning of the Camp Fire, and on the critical days of Pioneer and LNU, winds exceeded the published "
  "limits of every drone in the agency fleet. Water drop and logistics support drones exist only as tests and demonstrators. The "
  "uncertainty model therefore carries drone availability by wind and time of day as an explicit input to the crew safety "
  "margin.")

# ======================================================================== 1
H("1. Purpose and the decision numbers")
P("NASA's ACERO ConOps sets the goal of 24/7 aerial operations, organised in this project as surveillance (M1 fireline "
  "mapping, M2 spot fire detection, M3 firefighter safety) and suppression (M4 water drop, M5 logistics support: water, hose, pumps and supplies to crews). Crewed "
  "aircraft can safely work about 30% of a 24-hour cycle [NASA ACERO ConOps 2026, pp.30-31], and nearly 20% of ignitions "
  "occur outside traditional flight hours, some growing uncontested for up to 16 hours [NASA ACERO ConOps 2026, pp.15-16]. "
  "This proposal turns the effect of 24/7 aerial operations on ground crews into two numbers with error bars:")
B([
 "Crew safety (M3): the safety margin M = T_fire - (T_detect + T_notify + T_egress), the time between fire arrival and a "
 "crew's completed escape, and the probability that it goes negative, P(M < 0).",
 "Logistics support (M5): the probability that a requested delivery arrives within its window, P(on time).",
])
P("P(on time) is built from the four M5 questions on the Week 6 M5 Logistics Support tab, unchanged: Q1, can the drone carry the required water, suppressant or water-handling equipment (hose spools, pumps, drinking water) to the fire area or to firefighters (useful payload); Q2, can it reach crews far from roads or staging areas and return safely with the payload (loaded range, endurance with payload, turnaround); Q3, can water and supplies be delivered accurately and safely where vehicle access is limited (drop or landing accuracy, delivery success rate); Q4, can the fleet provide enough capacity to be operationally useful (delivery time, deliveries per hour, throughput, number of drones).")
P("Research questions. RQ1: how does P(M < 0) change with and without 24/7 aerial surveillance? RQ2: what is P(on time) "
  "for aerial delivery by day, night and smoke, at real incident demand? RQ3: which inputs dominate the uncertainty? RQ4: how "
  "much does fresher observation shrink it, against the measured baseline of Sections 3 and 4? RQ5: how large and biased are "
  "the operational records themselves, so that reporting error is not mistaken for model error?")

# ======================================================================== 2
H("2. What happened on the four fires")
H("2.1 The fire versus the record", 2)
FIGURE(HERE / "figures" / "fig1_size_on_record_vs_fire.png",
       "Figure 1. Fire size on record against the fire. Orange: ICS-209 size as reported (CURR_INCIDENT_AREA) at the end of each "
       "report period. Blue: aircraft perimeters (NIFC GeoMAC). Green: satellite area (NASA FEDS 12-hourly VIIRS; GOFER hourly "
       "GOES for LNU; the dips on LNU are reconstruction artifacts). Purple bands: documented crew hazard windows (Camp: 19 NIST "
       "burnovers; Carr: 26 July evening). Satellite area reads high on Carr and LNU (Section 4).")
P("The pattern is the same on all four fires and the gap is widest exactly when growth is fastest. On Pioneer, a long "
  "timber fire, the record kept within about a day of the perimeter. On Carr, Camp and LNU, wind-driven runs put the fire "
  "tens of thousands of acres ahead of the official size for one to two days, and the crew hazards on Carr and Camp fall "
  "inside those windows.")

H("2.2 Fire by fire", 2)
P("Pioneer 2016 (Boise NF, Idaho; 188,404 acres). Reported by the Casner Fire lead plane at 17:17 on 18 July at 1 to 1.5 "
  "acres, about 100 acres by 21:27; poor access and no nearby helicopter landing site complicated initial attack "
  "(Gabbert 2016, perimeter supplement). The first perimeter came 13.7 hours after the first report. The fire grew slowly "
  "for a month, then ran about 29,000 acres to the northeast on 29 August (Rodriguez et al. 2020, supplement) with "
  "pyroconvective updrafts [Rodriguez et al 2020]. 91% of its final area grew while the newest perimeter was more than 12 hours "
  "old, and the archive has a 140-hour gap with no perimeter in late September. No firefighter entrapment was found.")
P("Carr 2018 (Shasta and Whiskeytown, California; 229,651 acres). Ignition 13:09 on 23 July, first calls 13:10 to 13:11, air "
  "attack estimated 20 acres at 13:32 (USFS origin investigation). Satellites saw it within about an hour; the first "
  "reliable aircraft perimeter came 14.2 hours after ignition. On the evening of 26 July the fire made a wind-driven run with a "
  "fire whirl: the air tactical supervisor warned Branch III to leave at 18:38 and a mayday followed at 19:40, 62 minutes later "
  "(Carr Green Sheet). The newest ICS-209 at that moment was 10.5 hours old and gave 20,000 acres; the satellite showed about "
  "57,700 and the perimeter 44,484. The next day's ICS-209 still gave 48,312 acres against a perimeter of 80,968.")
P("Camp 2018 (Butte County, California; 153,336 acres, 85 deaths). First 911 call 06:25 on 8 November, first engine in sight "
  "at 06:44 reporting about 35 mph sustained wind [Maranghides et al 2023, pp.87-88; Maranghides et al 2021, p.91]. The fire "
  "reached Paradise by 08:01 and the first burnovers began at 07:50, 85 minutes after the first call. The first ICS-209 was "
  "filed at 18:30 for 16,000 acres. Section 2.3 gives the first 36 hours.")
P("LNU Lightning Complex 2020 (Napa, Sonoma, Lake, Solano, Yolo; 363,220 acres). Over 2,000 lightning strikes on 16 to "
  "17 August started multiple fires; start times of 06:00, 06:40 and 08:00 appear in different records (supplement). Fires merged overnight on 18 to "
  "19 August with a 12-mile run while resources were limited (CAL FIRE and Solano OES 2021, supplement). The complex's first "
  "ICS-209 in ICS-209-PLUS was filed 50 hours after the reference start, already at 33,025 acres. Half the final area had burned by "
  "the fourth day and 37% of growth happened between 20:00 and 06:00. No firefighter entrapment was found.")

H("2.3 Camp 2018: the first 36 hours from ICS-209-PLUS", 2)
P("ICS-209-PLUS holds no perimeter geometry. For Camp it gives 38 reports, each with a report period, the size as reported, "
  "percent contained, fire behaviour and narrative, plus an incident-level link to the FIRED satellite fire event. Those "
  "fields were combined with the GeoMAC and FEDS histories and the NIST and CAL FIRE timelines into a 128-row time-area table "
  "in the same layout as the team's supplement (perimeter_supplements/Camp_2018_Perimeter_Time_Area_ICS209PLUS.xlsx).")
FIGURE(HERE / "figures" / "fig2_camp_first_36h.png",
       "Figure 2. Camp Fire, first 36 hours. What was on record (ICS-209, orange) against aircraft perimeters, satellite area "
       "and radio estimates. Purple bands: burnover windows (Maranghides et al 2021, p.197; CAL FIRE Green Sheet). The first GeoMAC "
       "perimeter is stamped 09:54; its size matches the Green Sheet's first-12-hour estimate, so the stamp may be a processing "
       "time rather than the flight time.")
T(["Time, 8 Nov", "What happened", "What was on record"], [
 ["06:25 to 06:44", "First 911 call 06:25; dispatch 06:31; first engine in sight 06:44, about 35 mph wind", "Nothing; the ICS-209 discovery time is later set to 08:29"],
 ["06:55 to 07:37", "Pulga evacuation order 06:55 (3 min before fire arrived); engine reports 200 to 300 acres at 07:10; Concow order 07:37", "Radio only"],
 ["07:50 to 08:01", "First burnovers (Hoffman Road, Concow Road) 07:50; air attack about 1,000 acres at 07:52 and 'now in Paradise, about 2,000 acres' at 08:01", "Radio only"],
 ["08:30 to 13:10", "Burnovers 3 to 14 in Paradise and Concow; nearly all of Paradise and Magalia under orders by 09:03", "First aircraft perimeter (stamped 09:54, 54,586 acres); first satellite detection 10:14 (MODIS)"],
 ["14:15 to 15:15", "Rattlesnake Flats: firing operation planned 14:15; wind shift 14:45; burnover logged 15:15; an emergency call went unanswered because of radio traffic", "Newest perimeter 4.3 to 5.4 h old; no ICS-209 yet"],
 ["18:30", "CAL FIRE later estimates more than 55,000 acres in the first 12 h", "First ICS-209: 16,000 acres (18,000 at 19:20)"],
 ["9 Nov 07:45", "", "ICS-209: 70,000 acres. ICS-209-PLUS dates the fire's largest growth (52,000 acres) to this report, a day after the run"],
], widths=[1.0, 3.2, 2.5])
P("Two lessons for the uncertainty model. First, the derived growth fields in ICS-209-PLUS (largest growth, growth date) "
  "inherit the reporting lag, so they cannot time a blowup. Second, the linked FIRED event gives 117,968 acres, 23% below the "
  "official final size, and places ignition about 17 km west of the origin, at the first burned MODIS pixel; satellite "
  "products need a per-fire bias term (U14), and their ignition points are not detection times.")

# ======================================================================== 3
H("3. How fast the first information got out")
P("The safety margin starts with detection and notification. The table follows the first information on each fire from the "
  "first human report to the first entry in each record. Satellite times are acquisition times from the FIRMS archive "
  "(first overpass with a detection within 15 km of the origin, 30 km for LNU); delivery to users adds latency, up to 4 "
  "hours for VIIRS [Kochanski et al 2023, p.12].")
c = CH
T(["", "Pioneer 2016", "Carr 2018", "Camp 2018", "LNU 2020"], [
 ["First human report", "17:17 lead plane", "13:10 to 13:11 calls", "06:25 911 call", "Lightning, many ignitions"],
 ["First aircraft estimate", "17:17, 1 to 1.5 ac", "13:32, 20 ac (+0.4 h)", "07:52 radio, ~1,000 ac (+1.5 h)", "not in project records"],
 ["First satellite detection", "MODIS +6.3 h", "MODIS +1.0 h", "MODIS +3.8 h", "MODIS +17.1 h (VIIRS not downloaded)"],
 ["First satellite area", "none (FEDS starts 2018)", f"FEDS +{c['Carr 2018']['h_to_first_satellite_area']} h", f"FEDS +{c['Camp 2018']['h_to_first_satellite_area']} h", f"GOFER hourly +{c['LNU 2020']['h_to_first_satellite_area']} h"],
 ["First aircraft perimeter", f"+{c['Pioneer 2016']['h_to_first_aircraft_perimeter']} h", f"+{c['Carr 2018']['h_to_first_aircraft_perimeter']} h", f"+{c['Camp 2018']['h_to_first_aircraft_perimeter']} h (stamp uncertain)", "no aircraft series"],
 ["First ICS-209 report", f"+{c['Pioneer 2016']['h_to_first_ICS209']} h, 100 ac", f"+{c['Carr 2018']['h_to_first_ICS209']} h, 1,500 ac", f"+{c['Camp 2018']['h_to_first_ICS209']} h, 16,000 ac", f"+{c['LNU 2020']['h_to_first_ICS209']} h, 33,025 ac (complex)"],
 ["ICS-209 discovery time", "+0.9 h", "+2.1 h", "+2.1 h", "+8.0 h"],
], widths=[1.4, 1.3, 1.3, 1.4, 1.4])
P("Hours are from the earliest documented start: Pioneer 17:17, Carr 13:09, Camp 06:25, LNU 06:00 (GOFER reference). A "
  "single VIIRS pixel 6 minutes after the Carr ignition sits inside the search radius but is not confirmed as this fire.", italic=True, size=9)
P("Reading. People detect fires within minutes; records lag by hours. The first aircraft estimate came within half an hour "
  "to 90 minutes where one is documented, but it lived on the radio. The first mapped perimeter took 3.5 to 14 hours and the "
  "first ICS-209 5 to 12 hours (50 hours for the LNU complex, whose component fires may have reported separately before the "
  "complex was formed). The ICS-209 discovery time is itself 1 to 8 hours after the first report, so it should never be "
  "used as T_detect (U17).")
P("Notification is the second link. On Camp, between 07:30 and 08:40 evacuation orders went out at about the same time "
  "the fire arrived in each zone, and in Concow the fire arrived before the warning [Maranghides et al 2023, pp.97-99]; the "
  "orders preceded the early burnovers by only 13 to 59 minutes [Maranghides et al 2023, Table 12, pp.100-101]. NIST notes that spread information was coming in from 911 calls because firefighters were busy saving lives "
  "[Maranghides et al 2021, pp.81-82]. At Rattlesnake Flats an emergency call went unanswered because of radio traffic "
  "(Camp Green Sheet). On Carr, 62 minutes passed between the aerial warning and the mayday. These set the priors for "
  "T_notify (U9) and show that a notification can fail outright, not only arrive late. Crew travel time (T_egress) uses "
  "published GPS travel rates (Sullivan 2020; Campbell 2017, 2019; research_data).")

# ======================================================================== 4
H("4. When information arrived versus when the fire grew")
P("The v0.4 evaluation measured the gap between the record and the fire on all four incidents. The key results, which "
  "become the baseline that drones have to beat (RQ4):")
T(["Measure", "Pioneer", "Carr", "Camp", "LNU", "Record"], [
 ["ICS-209 size equals previous night's perimeter", "48 of 56 days", "20 of 28", "4 of 13", "no GeoMAC", "ICS-209-PLUS, GeoMAC"],
 ["National report repeats previous evening's 209", "71 of 72 days", "37 of 38", "6 of 10", "28 of 28", "NICC IMSR"],
 ["Worst same-day perimeter / ICS-209 ratio", "1.12", "1.68 (27 Jul)", "3.91 (8 Nov)", "2.07 (satellite)", "GeoMAC, FEDS"],
 ["Median hours between aircraft perimeters (max)", "18.8 (140)", "17.9 (76)", "14.5 (93)", "none", "GeoMAC"],
 ["Share of final area grown while newest perimeter > 12 h old", "91%", "59%", "71%", "0% (hourly GOFER)", "gap analysis"],
 ["Share of aircraft perimeters with a night time stamp", "21%", "10%", "4%", "n/a", "GeoMAC"],
 ["Top three growth days, share of final area", "32%", "34%", "65%", "55%", "ICS-209-PLUS"],
], widths=[2.2, 0.8, 0.8, 0.8, 0.9, 1.2])
P("On the hourly LNU record, sampling the fire at a fixed revisit interval and comparing the last-known size with the "
  "reference every hour gives the cost of slow revisit (growth phase, 17 to 21 August):")
T(["Revisit interval", "2 h", "3 h", "6 h", "12 h", "24 h"], [
 ["Mean shortfall of last-known size", "2.6%", "5.1%", "10.9%", "20.0%", "32.2%"],
 ["90th percentile", "5.6%", "16.1%", "39.5%", "62.3%", "86.4%"],
], widths=[2.0, 0.8, 0.8, 0.8, 0.8, 0.8])
P("This sets the surveillance requirement used in Section 5: a revisit of 2 to 3 hours, around the clock, from the first "
  "hours of the fire. It measures how fast the size goes stale, not shape accuracy, and it comes from one fire with 2 km "
  "GOES pixels.")
P("Records also carry their own errors, which the model treats as an observation-error layer separate from model error "
  "(RQ5): the ICS-209-PLUS cleaned ACRES field is capped at the final size and hid LNU's 375,209 acre peak (use "
  "CURR_INCIDENT_AREA); FEDS area is biased per fire (+9% Carr, -2% Camp); two Carr GeoMAC perimeters are stamped before "
  "the ignition; personnel counts differ by a factor of 2 to 4 between NICC and ICS-209-PLUS. The full register is in the "
  "workbook (Data Quality sheet).")

# ======================================================================== 5
H("5. Where drones can close the gap")
H("5.1 Drone capability today and coming", 2)
P("Status classes: Flown with ACERO (in the library [Wu et al 2025, p.9]); In service (agency fleet, Week 5); Tested "
  "(flown on a fire test or competition, not approved for operations); In development. Week 5 figures are agency or "
  "vendor statements and are labelled W5.")
T(["Drone or class", "Status", "Endurance", "Payload", "Night", "Wind limit", "Fits missions"], [
 ["Alta X (Freefly), NWCG Type 3 rotorcraft", "Flown with ACERO; In service (USFS 22)", "50 min empty, 22 min at 22 lb (W5); FY22 average sortie 17 min (W5)", "35 lb max; IGNIS ignition", "Not shown in library", "22 mph (secondary, W5)", "M1, M2 close in; M3 checks"],
 ["SwitchBlade-Elite, Type 3 band", "In service (USFS 66, largest type, W5)", "Up to 50 min; FY22 average sortie 18 min (W5)", "2 to 2.3 kg", "not compiled", "25 mph (W5)", "M1, M2 close in"],
 ["DOI small UAS fleet, Type 3 and 4", "In service (531 aircraft FY2018; 10,342 flights FY2018, W5)", "Type 3: 20 to 60 min; Type 4: up to 30 min [NWCG 2025, p.5]", "EO / IR, ignition", "Night flights need a DOI/FAA agreement [NWCG 2025, pp.6-7]", "not published", "M1, M2 on or near the line"],
 ["L3Harris FVR-90 (Overwatch Aero), NWCG Type 1 fixed wing", "Flown with ACERO as relay; In service on USFS and BLM fires (W5)", "12 to 18 h (W5); Type 1 class 6 to 14 h [NWCG 2025, p.5]", "8 to 22 lb thermal video, radio relay", "Day and night (operator, W5); ACERO flights were VLOS", "not found", "M1, M2 wide area; M3 overwatch and relay"],
 ["SuperVolo (NASA Langley), VTOL fixed wing", "Flown with ACERO as alternate relay", "8 h (W5)", "Radio relay", "Day only in tests", "not found", "Relay; candidate mapping"],
 ["Type 1 and 2 contract fixed wing (ScanEagle, Aerosonde, Penguin C)", "In service by contract [NWCG 2025, p.5]", "Type 1: 6 to 14 h; Type 2: 1 to 6 h", "Gimbaled EO / IR", "Above all other aircraft; night by agreement", "not published", "M1 incident mapping; M2 wide-area IR"],
 ["Joby optionally piloted aircraft", "Flown with ACERO (high orbit, telemetry only)", "not found", "not found", "not shown", "not found", "Future night water drop class [NASA ACERO ConOps 2026, p.19]"],
 ["Ondas Optimus drone-in-a-box", "Available; wildfire use not found (W5)", "not found", "Sensor", "Automated launch from a station", "not found", "M2 first detection, if a station network exists"],
 ["Dryad Silvanet sensors plus Silvaguard drones", "Tested (XPRIZE finals, Alaska, 2026; W5)", "not found", "Up to 100 L suppressant (source disagrees)", "Sensors 24/7", "not found", "M2 detection in minutes; M4 small fires"],
 ["Data Blanket swarm", "Tested; FAA approval for one pilot flying many drones (W5)", "not found", "Thermal", "not found", "not found", "M1, M3 at scale with fewer pilots"],
 ["Seneca Argo-1 swarm (5 drones)", "Tested by CAL FIRE, July 2026; approvals pending (W5)", "not found", "20 gal water plus foam each; 500 to 1,000 gal foam per pack", "IR targeting", "not found", "M4 initial attack"],
 ["FireSwarm Thunder Wasp GT", "Tested (XPRIZE finalist; W5)", "30 to 90 min (sources disagree)", "Up to 100 gal water", "not found", "not found", "M4; M5 bulk water to crews"],
 ["Rain and Sikorsky autonomous Black Hawk / Firehawk", "Demonstrator; CAL FIRE agreement 2025 (W5)", "helicopter", "324 gal bucket (demo); 1,000 gal tank class", "Aimed at night and smoke", "Flew hands-off in gusts to 30 kt (W5)", "M4 at night; M5 heavy logistics"],
 ["Piasecki KARGO II; Elroy Air Chaparral", "In development (W5)", "not found", "1,500 lb; 500 lb over 300 nm", "not found", "not found", "M5 logistics support (crew resupply)"],
], size=7.5, widths=[1.35, 1.1, 1.05, 0.95, 0.85, 0.7, 0.95])
P("Two facts frame everything below. Agency drones today fly short sorties: the USFS fleet averaged 17 to 18 minutes per "
  "flight in FY22 (W5). And no drone in the library has flown water drops or logistics deliveries to crews; Week 5 adds only tests and "
  "demonstrators for M4 and M5 [NASA ACERO ConOps 2026, p.19; Boroujeni et al 2024, p.38].")

H("5.2 Gap by gap", 2)
T(["Gap on the four fires", "What a drone must do", "Current drones that fit", "Future drones (W5)", "What still blocks it"], [
 ["First hours. First perimeter 3.5 to 14 h after the first report; first ICS-209 5 to 12 h (Section 3). Camp burnovers began 85 min after the 911 call.",
  "Launch within the first hour, map and send a perimeter, keep revisiting",
  "Type 3 rotorcraft from an engine or station for a first look (20 to 60 min, 5 mi); Type 1 fixed wing (FVR-90) for the full fire, VTOL so no runway",
  "Drone-in-a-box stations (Ondas); sensor networks that trigger a drone (Dryad, tested); tower plus drone systems (Anduril, tested)",
  "Dispatch and airspace: NOTAM 'as soon as practical' in initial attack [NWCG 2025, pp.6-7]; a TFR is assumed before UAS work in smoke or at night [NASA ACERO ConOps 2026, p.8]"],
 ["Revisit. Aircraft perimeters 14.5 to 18.8 h apart (median), gaps to 140 h; 59% to 91% of growth unseen for more than 12 h. Requirement: 2 to 3 h.",
  "Hold a station over the fire and remap every 2 to 3 h, day and night",
  "Type 1 fixed wing: 6 to 18 h endurance, above all other aircraft [NWCG 2025, p.5]; FVR-90 already maps perimeters on fires (W5)",
  "Swarms with one pilot for many drones (Data Blanket)",
  "Pilots and airframes for continuous cover (5.3); perimeter length and sensor swath per fire not yet set"],
 ["Night. 37% of LNU growth between 20:00 and 06:00; 4% to 21% of perimeters stamped at night; crewed window about 30% of the day.",
  "Fly and map at night and in smoke (DVE)",
  "Drones have mapped perimeters at night [NASA ACERO ConOps 2026, p.14]; FVR-90 day and night (W5). ACERO TCL 1 flights were daytime VLOS [Wu et al 2025, p.9]",
  "Autonomous helicopters aimed at night work (Rain and Sikorsky)",
  "Night flight needs a DOI/FAA agreement [NWCG 2025, pp.6-7]; night BVLOS not demonstrated in the library"],
 ["Warning. Camp evacuation orders went out about when the fire arrived (07:30 to 08:40), 13 to 59 min before burnovers; an emergency call went unanswered; Carr warning to mayday 62 min.",
  "Watch the crew and its escape route, relay warnings when radios saturate",
  "FVR-90 and SuperVolo flew as airborne relays in TCL 1 [Wu et al 2025, p.9]; Type 3 checks of escape routes",
  "not found in W5",
  "Relay link: telemetry 52% success with four ground nodes, over 90% with three; a single relay does not scale [Wu et al 2025, pp.9-11]. No drone has demonstrated crew tracking or alert delivery"],
 ["Wind. Camp first morning 20 to 30 mph sustained, gusts over 40; Pioneer gusts 28 to 40; LNU 20 to 40; Carr up to 21 (Week 6 wind table).",
  "Fly through the wind event that drives the blowup",
  "Only Carr is within the Alta X limit (marginal). No published wind limit for Type 1 fixed wing in project records",
  "Autonomous Black Hawk flew in gusts to 30 kt (W5)",
  "The worst days are the windiest days. On Camp the aircraft were grounded and air attack was 'not very effective due to smoke and wind' [Maranghides et al 2021, p.403]"],
 ["Water drop (M4) and logistics support (M5). No library drone has flown either. On Camp, crews defending structures needed water and hose while roads were jammed by evacuation traffic (Week 6 Fire Cross-Walk).",
  "Drop water at night; carry supplies to crews on a schedule",
  "None",
  "Seneca Argo-1 (tested), FireSwarm, Dryad (XPRIZE), autonomous Firehawk (demo); KARGO II, Chaparral (in development)",
  "Payload per sortie is small (20 to 100 gal) against helicopters (300 to 1,000 gal); approvals pending"],
], size=7.5, widths=[1.5, 1.05, 1.5, 1.2, 1.6])

H("5.3 What 24/7 surveillance takes", 2)
P("For one continuous station over a fire (one aircraft always on station), the arithmetic below uses endurance from the "
  "table above and the DOI rule that remote pilots fly at most 8 hours per duty day (OPM-11, W5). Transit, turnaround and "
  "spares are not included, so these are lower bounds.")
T(["Platform", "Endurance used", "Sorties per 24 h", "Remote pilot shifts per 24 h"], [
 ["FVR-90 (Type 1)", "12 h", "2", "3"],
 ["Type 1 class, low end", "6 h", "4", "3"],
 ["SuperVolo", "8 h", "3", "3"],
 ["Alta X (Type 3)", "17 min average sortie (FY22)", "about 85", "3 per launch site, each flying about 28 sorties"],
], widths=[1.8, 1.8, 1.3, 2.0])
P("A Type 1 fixed wing can hold the 2 to 3 hour revisit with a handful of sorties a day; Type 3 rotorcraft cannot hold "
  "persistent cover and belong in the close-in tier (spot fire confirmation, escape route checks), as in the fleet mix on "
  "the Week 6 Platform tab. How many stations a fire needs depends on perimeter length and sensor swath, which is left to "
  "the platform analysis.")

H("5.4 What this means for the proposal", 2)
B([
 "Drones can, in principle, close the information gaps that dominate the first hours, the revisit cadence and the night "
 "(Sections 3 and 4) with aircraft classes already flying on fires. These enter the model as faster T_detect and fresher "
 "information age.",
 "Wind is the binding constraint, and it bites on the days that matter: Camp's first morning, Pioneer's and LNU's runs. "
 "Drone availability is therefore an uncertain input conditioned on wind (U21), not an assumption.",
 "Warning delivery (T_notify) is the weakest link: relays have flown, but alert delivery to crews has not been shown and link "
 "success was 52% in the field test.",
 "Water drop (M4) and logistics support (M5) remain future capability; they are evaluated as scenarios with Week 5 payloads "
 "against the M5 questions (payload, loaded range, accuracy, deliveries per hour), not as "
 "validated performance.",
])

# ======================================================================== 6
H("6. Uncertainty model")
P("Inputs carried from v0.4 (U1 to U20) are listed in the workbook (Uncertainty Register). The ones that matter most for "
  "the questions above, and the four added in v0.5:")
T(["ID", "Input", "Prior or range", "Source"], [
 ["U1", "Rate of spread model error", "MAPE 47%; timber underpredicts", "[Cardil et al 2023, pp.6-7]"],
 ["U5, U12", "Information age at decision", "Official size about 1 day old; perimeters 14.5 to 18.8 h apart", "Section 4"],
 ["U9", "Notification latency and failure", "Camp: orders about at fire arrival, 13 to 59 min before burnovers; unanswered call; Carr 62 min", "Section 3"],
 ["U8", "Crew egress time", "GPS travel-rate models", "Sullivan 2020; Campbell 2017, 2019"],
 ["U17", "Start and detection time", "ICS-209 discovery 1 to 8 h after the first report", "Section 3"],
 ["U18", "Last-known size shortfall by revisit", "2.6% at 2 h to 32.2% at 24 h (mean, growth phase)", "Section 4, LNU GOFER"],
 ["U21 (new)", "Drone availability by wind", "Fly / no fly by airframe limit against hourly wind; Type 1 limit unknown (sensitivity 25 to 40 mph)", "Week 5 and 6 Platform tabs; ASOS and ERA5 hourly (research_data)"],
 ["U22 (new)", "Time to first drone perimeter", "Launch plus transit plus mapping; first look 20 to 60 min (Type 3) as a range to elicit", "[NWCG 2025, p.5]; elicitation"],
 ["U23 (new)", "Night and BVLOS authorisation", "Probability a night or BVLOS mission is approved in time", "[NWCG 2025, pp.6-7]; elicitation"],
 ["U24 (new)", "Warning relay success", "52% to over 90% message success by number of ground nodes", "[Wu et al 2025, pp.9-11]"],
], widths=[0.8, 1.6, 2.6, 1.7])
P("Methods (unchanged from v0.4 in substance): ensembles of ELMFIRE or GridFire with spatially correlated perturbations of "
  "fuel, wind and moisture [Pyregence Consortium n.d. GridFire, p.113; Altintas 2025, p.8]; a model-error layer from Cardil's "
  "residuals and a separate observation-error layer from Section 4; perimeter assimilation as the lever for RQ4 "
  "[Kochanski et al 2023, p.12; Boroujeni et al 2024, p.46]; a discrete-event simulation for delivery sorties driven by "
  "recorded daily crew counts; PERT elicitation for terms with no data [Hope et al 2024, pp.7-8]; one-at-a-time then Sobol "
  "sensitivity analysis.")

# ======================================================================== 7
H("7. Experiment and validation")
T(["Factor", "Levels", "Basis"], [
 ["Fire case", "Camp 8 Nov and Carr 26 Jul (validation); Pioneer (long timber); LNU (complex, night growth)", "Section 2"],
 ["Information source", "Official record (about 1 day); aircraft perimeters (14 to 19 h); satellite 12 h; drone revisit 6, 3, 2, 1 h", "Sections 3 and 4"],
 ["Drone availability", "Day only; 24 h; 24 h with wind limit (U21); 24 h with wind limit and authorisation delay (U23)", "Section 5"],
 ["Time of day", "Day VMC; night DVE; smoke DVE", "[NASA ACERO ConOps 2026, pp.30-31]"],
 ["Warning relay", "52%; 87%; 95% message success", "[Wu et al 2025, pp.9-10]"],
 ["Model", "ELMFIRE; GridFire", "[Altintas 2025, p.8]"],
], widths=[1.3, 3.7, 1.6])
P("Run Camp 8 November and Carr 26 July first, because both have timed crew hazards. Validation: Camp's minute-level NIST "
  "progression and 19 burnovers [Maranghides et al 2021, p.148] and NETTRA notification times [Maranghides et al 2023, "
  "pp.97-101]; Carr's timed warning and mayday plus 63 GeoMAC and 88 FEDS perimeters. Pioneer and LNU validate spread and "
  "information age, not P(M < 0). Forecasts are scored against time-matched perimeters with the observation-error layer "
  "applied. A fractional design or Latin hypercube keeps the run count manageable.")

# ======================================================================== 8, 9
H("8. Limitations")
B([
 "Four fires; the figures describe these incidents, not a national rate.",
 "The revisit curve comes from one fire (LNU) and 2 km GOES pixels.",
 "Archive time stamps may record processing rather than flight time (Camp's first perimeter, two Carr perimeters).",
 "Week 5 drone figures are agency and vendor statements; Type 1 fixed-wing wind limits were not found.",
 "The FIRMS search is by distance from the origin and may include nearby unrelated fires in the first hour.",
 "Radio and dispatch logs are not in hand; notification times rest on published reports.",
 "No firefighter entrapment was found for Pioneer or LNU in one day of searching (30 September 2026).",
 "Novelty is not assessed; run the literature gate before any claim that this is new.",
])
H("9. Next steps")
N([
 "Team agrees on P(M < 0) and P(on time) as headline metrics, with information age and drone availability as reported covariates.",
 "Build U21: hourly wind for each fire's critical period from ASOS and ERA5 (research_data) against airframe limits; ask Overwatch Aero or NASA for an FVR-90 wind limit.",
 "Run 100-member ELMFIRE ensembles for Camp 8 November and Carr 26 July at the observed cadence and at 6, 3, 2 and 1 hour drone revisits, with and without the wind gate.",
 "Confirm the Camp first perimeter time (09:54 stamp) against NIST maps; ask whether GOFER covers Carr and Camp.",
 "Download FIRMS VIIRS for LNU and the component-fire ICS-209s to complete the first-report chain.",
 "When the Pioneer radio and dispatch logs arrive, add timed fire positions and notification times; consider records requests for Carr and Camp.",
 "Elicit U22 and U23 with the stakeholder questionnaire (first drone on scene, night approval time).",
 "Request NIROPS access for night infrared perimeters.",
])

H("References (project library)")
B([
 "Altintas 2025. Actionable fire modeling in Firemap for extended attack decision support.",
 "Boroujeni et al 2024. A comprehensive survey of research towards AI-enabled unmanned aerial systems in pre-, active-, and post-wildfire management.",
 "CAL FIRE 2018. Camp Fire Green Sheet (serious injury report, Rattlesnake Flats).",
 "Cardil et al 2023. Performance of operational fire spread models in California. International Journal of Wildland Fire.",
 "Hope et al 2024. A cost-benefit analysis of WildFireSat. PLOS ONE.",
 "Kochanski et al 2023. Analysis of methods for assimilating fire perimeters into a coupled fire-atmosphere model.",
 "Maranghides et al 2021. A case study of the Camp Fire: Fire progression timeline. NIST TN 2135.",
 "Maranghides et al 2023. A case study of the Camp Fire: Notification, evacuation, traffic, and temporary refuge areas (NETTRA). NIST TN 2252.",
 "NASA ACERO ConOps 2026. Kopardekar et al. 24/7 Wildland Fire Aerial Operations Airspace Management System ConOps.",
 "NWCG 2025. Standards for fire unmanned aircraft systems operations.",
 "Pyregence Consortium n.d. GridFire documentation.",
 "Rodriguez et al 2020. Extreme pyroconvective updrafts during a megafire.",
 "Wu et al 2025. Development of a portable airspace management system for wildland fire operations.",
])
H("Data sources")
B([
 "ICS-209-PLUS v3 (St. Denis et al. 2023), sitreps and incident tables; NICC Incident Management Situation Reports 2016, 2018, 2020.",
 "NIFC GeoMAC perimeter histories; NASA FEDS 12-hourly VIIRS perimeters; GOFER hourly GOES reconstruction (Liu et al. 2024); NASA FIRMS MODIS and VIIRS archive.",
 "Team perimeter data supplement (Pioneer, Carr, LNU) and the new Camp table (perimeter_supplements/).",
 "Research Data Index (timing anchors, Camp burnovers), Carr and Camp Green Sheets (research_data/entrapments).",
 "Weekly Spreadsheets: NASA_week6_updated.xlsx Platform tab (library drones, fleet mix, wind table) and NASA_week5.xlsx Platform tab (agency fleet and future aircraft, with source links in its Source column).",
])
P("Page numbers in library citations are PDF page numbers as indexed by the project RAG. Analysis: case_study_evaluation.py, "
  "perimeter_gap_evaluation.py, first_report_evaluation.py, build_uq_v05.py.", italic=True, size=9)

docx_path = HERE / "UQ_Crew_Safety_Logistics_Proposal_v0.5.docx"
d.save(docx_path)

# ======================================================================== workbook: v0.4 plus new sheets
src = HERE / "UQ_Evaluation_Workbook_v0.4.xlsx"; dst = HERE / "UQ_Evaluation_Workbook_v0.5.xlsx"
shutil.copy(src, dst)
wb = load_workbook(dst)
HDR = PatternFill("solid", fgColor="1F3864"); NEW = PatternFill("solid", fgColor="E2EFDA")
def sheet(title, header, rws, widths):
    if title in wb.sheetnames: del wb[title]
    ws = wb.create_sheet(title); ws.append(header)
    for r in rws: ws.append(r)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = HDR
    for row in ws.iter_rows():
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"; return ws

rd = wb["README"]
rd["B2"] = (f"v0.5, {DATE}. Adds sheets First Report Chain, Drone Gap Coverage, Drone Sortie Math, and Uncertainty Register rows "
            "U21 to U24 (green). M5 stays Logistics support; its Week 6 questions are on the M5 Questions (Week 6) sheet. Camp ICS-209-PLUS time-area table is a separate file in perimeter_supplements/. All v0.4 sheets kept.")
rd["B4"] = "case_study_evaluation.py, perimeter_gap_evaluation.py, first_report_evaluation.py, then build_uq_v05.py."

ur = wb["Uncertainty Register"]
for r in [
 ["U21", "Drone availability by wind", "M1, M2, M3", "Aleatory (weather)", "Hourly fly / no-fly from airframe limit vs observed wind", "Alta X 22 mph; SwitchBlade-Elite 25 mph; Type 1 fixed wing not found (sensitivity 25 to 40 mph). Camp, Pioneer, LNU critical periods exceed small-UAS limits", "n/a (gates the aerial observation)", "Week 6 Platform wind table; Week 5 Platform tab; ASOS / ERA5 hourly in research_data", "v0.5 new"],
 ["U22", "Time to first drone perimeter", "M1, M2, M3", "Epistemic", "Launch + transit + mapping time distribution", "Type 3: 20 to 60 min flight, 5 mi range; Type 1: VTOL, 50 mi range; elicit launch delay", "Yes", "[NWCG 2025, p.5]; elicitation", "v0.5 new"],
 ["U23", "Night and BVLOS authorisation", "M1, M2, M3", "Epistemic", "Probability approved within the needed time", "Night flights need a DOI/FAA MOA; NOTAM as soon as practical in initial attack; TFR assumed for DVE work", "No", "[NWCG 2025, pp.6-7]; [NASA ACERO ConOps 2026, p.8]", "v0.5 new"],
 ["U24", "Warning relay success", "M3", "Aleatory", "Bernoulli per message by ground node count", "Telemetry 52% (four ground nodes) to over 90% (three)", "Partly", "[Wu et al 2025, pp.9-11]", "v0.5 new"],
]:
    ur.append(r)
    for c in ur[ur.max_row]: c.fill = NEW; c.alignment = Alignment(wrap_text=True, vertical="top")

ch = list(csv.DictReader(open(HERE / "first_report_eval" / "first_report_chain.csv")))
sheet("First Report Chain", list(ch[0].keys()), [list(r.values()) for r in ch], [14, 16] + [30] * (len(ch[0]) - 2))

sheet("Drone Gap Coverage", ["Gap", "Measured on the four fires", "Requirement", "Current drones (library or W5 in service)", "Future drones (W5)", "Blocker", "Source"], [
 ["First hours", "First aircraft perimeter +3.5 to +14.2 h; first ICS-209 +5.3 to +12.1 h (LNU +50 h); Camp burnovers 85 min after first call", "Drone perimeter within the first hour", "Type 3 rotorcraft first look; Type 1 fixed wing VTOL (FVR-90)", "Drone-in-a-box (Ondas); Dryad sensors plus drones; Anduril towers plus drones", "NOTAM / TFR timing; dispatch", "first_report_eval; NWCG 2025 pp.5-7"],
 ["Revisit", "Median 14.5 to 18.8 h between perimeters; 59 to 91% of growth while newest perimeter > 12 h old", "2 to 3 h revisit, 24 h a day", "Type 1 fixed wing 6 to 18 h endurance", "One-to-many swarms (Data Blanket)", "Pilots and airframes; swath vs perimeter length not set", "perimeter_eval; Revisit Curve sheet"],
 ["Night", "37% of LNU growth 20:00 to 06:00; 4 to 21% of perimeters at night", "Night and DVE mapping", "FVR-90 day and night (operator claim)", "Autonomous Firehawk (Rain / Sikorsky)", "DOI/FAA night MOA; night BVLOS not in library", "ConOps p.14; NWCG pp.6-7"],
 ["Warning", "Camp orders about at fire arrival (13 to 59 min before burnovers); unanswered emergency call; Carr 62 min warning to mayday", "Crew overwatch and warning relay", "FVR-90, SuperVolo relays (TCL 1)", "not found", "Relay telemetry 52% to 90%; no alert delivery shown", "Wu et al 2025 pp.9-11; NETTRA Table 12"],
 ["Wind", "Camp 20 to 30 mph, gusts > 40; Pioneer gusts 28 to 40; LNU 20 to 40; Carr up to 21", "Fly through the blowup wind", "Alta X 22 mph (Carr only, marginal)", "Autonomous Black Hawk gusts to 30 kt", "Worst days are windiest; Type 1 limit unknown", "Week 6 wind table"],
 ["Water drop (M4) and logistics support (M5)", "No library drone has flown either", "Night water drop; scheduled delivery of water, hose, pumps and supplies to crews", "None", "Seneca Argo-1, FireSwarm, Dryad, autonomous Firehawk; KARGO II, Chaparral", "Small payload; approvals pending", "Week 5 Platform tab; ConOps p.19"],
], [14, 40, 22, 30, 30, 30, 26])

sheet("Drone Sortie Math", ["Platform", "Endurance used (h)", "Sorties per 24 h (one station)", "Remote pilot shifts (8 h flight / duty day)", "Notes"], [
 ["FVR-90 (Type 1)", 12, 2, 3, "Endurance 12 to 18 h (W5). Transit, turnaround and spares not included."],
 ["Type 1 class, low end", 6, 4, 3, "NWCG 2025 p.5 band 6 to 14 h."],
 ["SuperVolo", 8, 3, 3, "W5 / Wu et al 2025 p.9."],
 ["Alta X (Type 3)", round(17 / 60, 2), round(24 / (17 / 60)), "3 per site", "FY22 USFS average sortie 17.0 min (W5). Close-in tier only."],
], [24, 18, 22, 26, 60])

# M5 stays Logistics support; its questions are copied verbatim from the Week 6 M5 Logisitic Support tab
sheet("M5 Questions (Week 6)", ["Q#", "Week 6 question (unchanged)", "Week 6 metrics", "Feeds", "UQ inputs"], [
 ["Q1", "Can the drone carry the required water, suppressant, or water-handling equipment (hose spools, pumps, drinking water) to the fire area or to firefighters?",
  "Useful payload capacity", "Crews served per aerial asset", "U16 demand scale"],
 ["Q2", "Can it reach crews operating far from roads or staging areas and return safely with the required payload?",
  "Loaded mission range / radius; endurance with payload; turnaround time", "Turnaround p50 / p90; P(abort)", "U3 wind; U7 path deviation; U21 drone availability by wind"],
 ["Q3", "Can water and supplies be delivered accurately and safely where conventional vehicle access is limited?",
  "Landing / drop accuracy; delivery success rate", "P(on time); P(abort)", "U7 path deviation; U3 wind"],
 ["Q4", "Can the drone fleet provide enough water delivery capacity to make it operationally useful?",
  "Delivery time; deliveries per hour; water / supply throughput; number of drones required", "P(on time); crews served per aerial asset", "U10 delivery window; U16 demand scale; U11 human factors delay; U23 night and BVLOS authorisation"],
], [6, 60, 40, 30, 40])

for ws in wb:  # rule: no en or em dashes anywhere, including text carried over from v0.4
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and ("–" in c.value or "—" in c.value):
                c.value = c.value.replace("–", " to ").replace("—", ", ")
wb.save(dst)
print("wrote", docx_path.name, dst.name)
