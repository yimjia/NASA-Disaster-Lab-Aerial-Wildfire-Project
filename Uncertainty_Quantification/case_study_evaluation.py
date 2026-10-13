"""Evaluates the Fire Case Studies (Pioneer 2016, Carr 2018, Camp 2018, LNU 2020) for the UQ work.

Reads the case study workbooks (ICS-209 daily reports + NICC daily situation reports),
the GeoMAC perimeter histories and NASA FEDS 12 h perimeters in research_data/, and
writes case_study_eval/results.json plus CSV tables used by build_uq_v02.py.

Run: uv run --project ~/nasa_rag python case_study_evaluation.py
"""
import json, csv, statistics as st
from datetime import datetime, timedelta, timezone
from pathlib import Path
from openpyxl import load_workbook

HERE = Path(__file__).parent
ROOT = HERE.parent
CS = ROOT / "Fire Case Studies"
PER = HERE / "research_data" / "datasets" / "01_progression_perimeters"
OUT = HERE / "case_study_eval"; OUT.mkdir(exist_ok=True)
ACRES_PER_KM2 = 247.105

FIRES = {
    "Pioneer 2016": dict(wb=CS / "Pioneer_2016" / "Pioneer_Fire_2016_ICS209.xlsx", utc_off=-6,
                         geomac="Pioneer_2016_NIFC_GeoMAC_perimeter_history.geojson", feds=None, start="2016-07-18"),
    "Carr 2018": dict(wb=CS / "Carr_2018" / "Carr_Fire_2018_ICS209.xlsx", utc_off=-7,
                      geomac="Carr_2018_NIFC_GeoMAC_perimeter_history.geojson",
                      feds="Carr_2018_NASA_FEDS_perimeter_12h.geojson", start="2018-07-23"),
    "Camp 2018": dict(wb=CS / "Camp_2018" / "Camp_Fire_2018_ICS209.xlsx", utc_off=-8,
                      geomac="Camp_2018_NIFC_GeoMAC_perimeter_history.geojson",
                      feds="Camp_2018_NASA_FEDS_perimeter_12h.geojson", start="2018-11-08"),
    "LNU 2020": dict(wb=CS / "LNU_2020" / "LNU_Lightning_Complex_2020_ICS209.xlsx", utc_off=-7,
                     geomac=None, feds="LNU_2020_NASA_FEDS_perimeter_12h.geojson", start="2020-08-17"),
}

def num(v):
    try: return float(str(v).replace(",", ""))
    except (TypeError, ValueError): return None

def pct(a, b): return 100.0 * (a - b) / b if b else None

def q(xs, p):
    xs = sorted(xs)
    if not xs: return None
    k = (len(xs) - 1) * p; f = int(k); c = min(f + 1, len(xs) - 1)
    return xs[f] + (xs[c] - xs[f]) * (k - f)

def summ(xs):
    xs = [x for x in xs if x is not None]
    if not xs: return dict(n=0)
    return dict(n=len(xs), min=min(xs), p10=q(xs, .1), median=q(xs, .5), p90=q(xs, .9), max=max(xs), mean=st.mean(xs))

def sheet_rows(ws):
    rows = list(ws.iter_rows(values_only=True)); h = rows[0]
    return [dict(zip(h, r)) for r in rows[1:] if any(v is not None for v in r)]

results, tables = {}, {"size_by_day": [], "cadence": [], "resources": [], "growth": []}

for fire, cfg in FIRES.items():
    wb = load_workbook(cfg["wb"], data_only=True)
    R = {}
    # ---- ICS-209 reports (local time as recorded)
    # Use the as-reported size (CURR_INCIDENT_AREA). ICS-209-PLUS ACRES is cleaned and capped at the final size.
    r209 = [r for r in sheet_rows(wb["All 209 fields"]) if r.get("REPORT_TO_DATE")]
    for r in r209:
        r["t"] = datetime.fromisoformat(str(r["REPORT_TO_DATE"]))
        r["acres"] = num(r.get("CURR_INCIDENT_AREA")) or num(r["ACRES"]); r["acres_clean"] = num(r["ACRES"])
    r209.sort(key=lambda r: r["t"])
    gaps = [(b["t"] - a["t"]).total_seconds() / 3600 for a, b in zip(r209, r209[1:])]
    R["209_reports"] = len(r209)
    R["209_interval_h"] = summ(gaps)
    R["209_gaps_over_24h"] = sum(1 for g in gaps if g > 24)
    R["209_longest_gap"] = max(((b["t"] - a["t"]).total_seconds() / 3600, str(a["t"]), str(b["t"])) for a, b in zip(r209, r209[1:]))
    last209 = {}
    for r in r209:
        if r["acres"] is not None: last209[r["t"].date()] = r["acres"]
    dec209 = [(str(a), last209[a], last209[b]) for a, b in zip(sorted(last209), sorted(last209)[1:]) if last209[b] < last209[a]]
    R["209_downward_revisions"] = dec209
    R["209_rows_reported_ne_cleaned"] = sum(1 for r in r209 if r["acres"] != r["acres_clean"])
    R["209_max_reported_vs_cleaned"] = (max(r["acres"] for r in r209 if r["acres"]), max(r["acres_clean"] for r in r209 if r["acres_clean"]))

    # ---- NICC IMSR (publication date; numbers are as of the previous evening)
    nicc = {}
    for r in sheet_rows(wb["NICC Daily IMSR"]):
        d = r.get("Report date")
        if isinstance(d, str) and len(d) == 10 and isinstance(r.get("Size (acres)"), (int, float)):
            nicc[datetime.fromisoformat(d).date()] = r
    # lag test: does IMSR on day D match the last 209 on D-1 or on D?
    m_prev = m_same = n = 0; lag_err = []
    for d, r in nicc.items():
        a = r["Size (acres)"]
        if (d - timedelta(days=1)) in last209 and d in last209:
            n += 1
            if a == last209[d - timedelta(days=1)]: m_prev += 1
            if a == last209[d]: m_same += 1
        if (d - timedelta(days=1)) in last209:
            lag_err.append(pct(a, last209[d - timedelta(days=1)]))
    R["nicc_reports_with_row"] = len(nicc)
    R["nicc_lag_test"] = dict(days_compared=n, exact_match_prev_day_209=m_prev, exact_match_same_day_209=m_same)
    R["nicc_vs_prev_day_209_pct"] = summ([abs(x) for x in lag_err])

    # ---- GeoMAC perimeter history
    geo = []
    if cfg["geomac"]:
        for f in json.load(open(PER / cfg["geomac"]))["features"]:
            p = f["properties"]
            if p.get("perimeterdatetime") and p.get("gisacres"):
                t = datetime.fromtimestamp(p["perimeterdatetime"] / 1000, timezone.utc) + timedelta(hours=cfg["utc_off"])
                geo.append(dict(t=t.replace(tzinfo=None), acres=p["gisacres"], method=p.get("mapmethod")))
    geo.sort(key=lambda g: g["t"])
    if geo:
        gg = [(b["t"] - a["t"]).total_seconds() / 3600 for a, b in zip(geo, geo[1:])]
        R["geomac_perimeters"] = len(geo)
        R["geomac_interval_h"] = summ([g for g in gg if g > 0])
        R["geomac_share_interval_over_24h"] = sum(1 for g in gg if g > 24) / len(gg)
        R["geomac_methods"] = {m: sum(1 for g in geo if g["method"] == m) for m in sorted({g["method"] for g in geo}, key=str)}

    # ---- NASA FEDS 12 h satellite perimeters (sum of all fireids at each overpass)
    feds = {}
    if cfg["feds"]:
        for f in json.load(open(PER / cfg["feds"]))["features"]:
            p = f["properties"]; t = datetime.fromisoformat(p["t"])
            if t < datetime.fromisoformat(cfg["start"]): continue
            feds[t] = feds.get(t, 0) + p["farea"] * ACRES_PER_KM2
        R["feds_overpasses"] = len(feds)

    # ---- daily size table and cross-source differences (209 as reference)
    days = sorted(set(last209) | {d - timedelta(days=1) for d in nicc} | {g["t"].date() for g in geo} |
                  {(t + timedelta(hours=cfg["utc_off"])).date() for t in feds})
    diffs = {"NICC (prev-evening)": [], "GeoMAC": [], "FEDS": []}
    for d in days:
        g_d = [g["acres"] for g in geo if g["t"].date() == d]
        f_d = [a for t, a in feds.items() if (t + timedelta(hours=cfg["utc_off"])).date() == d]
        n_d = nicc.get(d + timedelta(days=1), {}).get("Size (acres)")
        row = dict(fire=fire, date=str(d), ics209_last=last209.get(d), nicc_next_morning=n_d,
                   geomac_last=g_d[-1] if g_d else None, feds_last=f_d[-1] if f_d else None)
        tables["size_by_day"].append(row)
        ref = last209.get(d)
        if ref and ref > 1000:
            if n_d: diffs["NICC (prev-evening)"].append(pct(n_d, ref))
            if g_d: diffs["GeoMAC"].append(pct(g_d[-1], ref))
            if f_d: diffs["FEDS"].append(pct(f_d[-1], ref))
    # latency chain: does the 209 on day D+1 equal the last GeoMAC perimeter of day D (within 0.5%)?
    if geo:
        gl = {}
        for g in geo: gl[g["t"].date()] = g["acres"]
        tested = [(d, gl[d], last209.get(d + timedelta(days=1)), last209.get(d)) for d in gl
                  if last209.get(d + timedelta(days=1)) and last209.get(d) and gl[d] > 1000 and last209.get(d + timedelta(days=1)) != last209.get(d)]
        nxt = sum(1 for d, a, b, c in tested if abs(b - a) / a < 0.005)
        same = sum(1 for d, a, b, c in tested if abs(c - a) / a < 0.005)
        R["latency_chain_geomac_to_209"] = dict(days_tested=len(tested), match_next_day_209=nxt, match_same_day_209=same,
                                                 note="days where the 209 size changed; match = within 0.5%")
    # blowup-day shortfall of the official 209 versus same-day perimeter / satellite
    short = []
    _ds = sorted(last209); _prev = {b: last209[a] for a, b in zip(_ds, _ds[1:])}; _prev[_ds[0]] = 0
    for d in sorted(last209, key=lambda d: -(last209[d] - _prev[d]))[:3]:
        g_d = [g["acres"] for g in geo if g["t"].date() == d]
        f_d = [a for t, a in feds.items() if (t + timedelta(hours=cfg["utc_off"])).date() == d]
        short.append(dict(date=str(d), ics209=last209[d], geomac=g_d[-1] if g_d else None, feds=f_d[-1] if f_d else None,
                          ratio_geomac=(g_d[-1] / last209[d]) if g_d else None, ratio_feds=(f_d[-1] / last209[d]) if f_d else None))
    R["blowup_day_shortfall"] = short
    R["size_diff_vs_209_pct"] = {k: dict(signed=summ(v), absolute=summ([abs(x) for x in v])) for k, v in diffs.items() if v}

    # ---- growth from the 209 last-report-per-day series
    ds = sorted(last209); grow = []
    for a, b in zip(ds, ds[1:]):
        if (b - a).days == 1:
            grow.append(dict(date=str(b), acres=last209[b] - last209[a], rel=(last209[b] - last209[a]) / last209[a] if last209[a] else None))
    pos = [g["acres"] for g in grow if g["acres"] > 0]; total = last209[max(last209)]
    top = sorted(grow, key=lambda g: -g["acres"])[:3]
    R["growth"] = dict(final_acres_209=total, days_with_growth=len(pos), daily_growth_acres=summ(pos),
                       top3_days=top, top3_share_of_final=sum(t["acres"] for t in top) / total,
                       days_to_50pct=next((i for i, d in enumerate(ds) if last209[d] >= total / 2), None),
                       days_to_90pct=next((i for i, d in enumerate(ds) if last209[d] >= 0.9 * total), None))
    for g in grow: tables["growth"].append(dict(fire=fire, **g))

    # ---- resources and cost from NICC (all assigned personnel, crews, engines, helicopters)
    nr = [r for r in nicc.values()]
    pers = [r["Personnel total"] for r in nr if isinstance(r.get("Personnel total"), (int, float)) and r["Personnel total"] > 0]
    def col(k): return [r[k] for r in nr if isinstance(r.get(k), (int, float))]
    R["resources_nicc"] = dict(personnel=summ(pers), personnel_days=sum(pers), crews=summ(col("Crews")),
                               engines=summ(col("Engines")), helicopters=summ(col("Helicopters")),
                               helicopter_days=sum(col("Helicopters")),
                               final_cost=max(col("Cost to date ($)") or [0]))
    ac = [r for r in nr if isinstance(r.get("Personnel total"), (int, float)) and r["Personnel total"] > 0 and isinstance(r.get("Helicopters"), (int, float))]
    R["resources_nicc"]["helicopters_per_1000_personnel"] = summ([1000 * r["Helicopters"] / r["Personnel total"] for r in ac])
    R["resources_nicc"]["crews_per_helicopter"] = summ([r["Crews"] / r["Helicopters"] for r in ac if r["Helicopters"] and isinstance(r.get("Crews"), (int, float))])
    for d, r in sorted(nicc.items()):
        tables["resources"].append(dict(fire=fire, imsr_date=str(d), personnel=r.get("Personnel total"), crews=r.get("Crews"),
                                        engines=r.get("Engines"), helicopters=r.get("Helicopters"), cost=r.get("Cost to date ($)"),
                                        pct_contained=r.get("% Ctn/Comp")))
    # 209 vs NICC personnel definitions
    p209 = [num(r.get("TOTAL_PERSONNEL")) for r in r209 if num(r.get("TOTAL_PERSONNEL"))]
    R["personnel_definition_gap"] = dict(nicc_peak=max(pers) if pers else None, ics209plus_peak=max(p209) if p209 else None,
                                         ratio=(max(pers) / max(p209)) if pers and p209 else None)

    # ---- final size across sources (summary sheet + perimeter sheet)
    fin = {"ICS-209 final": total}
    if nicc: fin["NICC IMSR last listed"] = nicc[max(nicc)]["Size (acres)"]
    if geo: fin["GeoMAC last perimeter"] = geo[-1]["acres"]
    if feds: fin["NASA FEDS last overpass"] = feds[max(feds)]
    ps = wb["Final Perimeter"]
    for r in sheet_rows(ps):
        v = r.get("GIS acres") if "GIS acres" in r else (r.get("Value") if r.get("Field") == "GIS_ACRES" else None)
        lab = r.get("Agency / source") or r.get("Fire") or ("NIFC final perimeter" if r.get("Field") == "GIS_ACRES" else None)
        if r.get("Fire") and not str(r.get("Fire")).startswith("TOTAL"): continue  # complex component, not whole fire
        if isinstance(v, (int, float)) and lab:
            fin[f"Perimeter: {lab}"] = v
    R["final_size_sources"] = fin
    vals = [v for v in fin.values() if isinstance(v, (int, float)) and v > 0.5 * total]
    R["final_size_spread_pct"] = 100 * (max(vals) - min(vals)) / total
    results[fire] = R
    for k in ("209_interval_h",):
        tables["cadence"].append(dict(fire=fire, source="ICS-209", **{x: R[k].get(x) for x in ("n", "median", "p90", "max")}))
    if geo: tables["cadence"].append(dict(fire=fire, source="GeoMAC perimeters", **{x: R["geomac_interval_h"].get(x) for x in ("n", "median", "p90", "max")}))
    if feds: tables["cadence"].append(dict(fire=fire, source="NASA FEDS (VIIRS)", n=len(feds), median=12, p90=12, max=12))

json.dump(results, open(OUT / "results.json", "w"), indent=1, default=str)
for name, rows in tables.items():
    if rows:
        with open(OUT / f"{name}.csv", "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=list(rows[0].keys())); w.writeheader(); w.writerows(rows)
print(json.dumps(results, indent=1, default=str))
